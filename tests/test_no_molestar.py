"""Pruebas unitarias y de integración de no_molestar.py (spec revisión 3; sin red ni navegador real).

Organización: contrato del bloque de configuración -> utilidades puras -> directorio SERNAC ->
entradas -> tracker (estructura real) -> mapeo -> correo -> BotSernac (FakePage + iframe) -> main().
"""
from __future__ import annotations

import builtins
import contextlib
import csv
import getpass
import io
import logging
import os
import re
import shutil
import smtplib
import socket
import ssl
import subprocess
import sys
from collections import Counter
from datetime import date, datetime
from email.message import EmailMessage
from email.utils import parseaddr
from pathlib import Path

import openpyxl
import pytest
from playwright.sync_api import Error as PWError
from playwright.sync_api import TimeoutError as PWTimeoutError

import no_molestar as nm
from conftest import (
    BBOX_NULO,
    CAJA_V2,
    CONFIG_VALIDA,
    CORREO_PERSONA,
    ENCABEZADOS_REALES,
    ENTRADA_BANCO_CHILE,
    ENTRADAS_CLARO,
    ENTRADAS_ENTEL,
    FORMATO_FECHA_REAL,
    GMAIL_USUARIO,
    JWT_PORTAL,
    PLANTILLA_VALIDA,
    RESULTADOS_WALK,
    SECRETO_CU,
    SECRETO_GMAIL,
    SECRETO_GMAIL_SIN_ESPACIOS,
    SELECTORES_QUE_VEN_EL_BADGE,
    TEL_PERSONA_E164,
    TEXTO_POR_ID,
    URL_ANCHOR_V2,
    URL_BFRAME,
    URL_HCAPTCHA,
    URL_TURNSTILE,
    BuscadorFalso,
    Elemento,
    FabricaSMTP,
    PortalFalso,
    PreguntarGuionado,
    S,
    SleepGrabador,
    _norm,
    buscar_en_directorio,
    celda,
    celdas,
    crear_tracker_real,
    encabezados_de,
    es_fecha,
    fila_de,
    fila_real,
    filas_reales,
    marcar_captcha,
    poner_badge_invisible,
    tiempos_prueba,
    todas_las_celdas,
    vacio,
    valor_celda,
)

HOY = date(2026, 10, 22)
RAIZ = Path(nm.__file__).resolve().parent
BOM = b"\xef\xbb\xbf"


def lanzar(exc: BaseException):
    def _f(*args, **kwargs):
        raise exc

    return _f


def _prohibido(nombre):
    def _f(*args, **kwargs):
        pytest.fail(f"{nombre} no debía llamarse", pytrace=False)

    return _f


# =============================================================================
# Contrato del bloque de configuración (estructura; no requiere implementación)
# =============================================================================
SELECTORES_ESPERADOS = {
    "login_boton_claveunica": ["#btnClaveUnica"],
    "cu_run": ["#uname"],
    "cu_clave": ["#pword"],
    "cu_ingresar": ["#login-submit"],
    "cu_error": [
        ".alert-danger",
        ".errorlist",
        "text=/(RUN|ClaveÚnica|clave).{0,40}(incorrect|inv[aá]lid|no coincide)|bloquead/i",
    ],
    "sesion_iniciada": ["button[onclick='logoutCU()']"],
    "boton_mantener_sesion": ["#sessionTimeoutModal #btnKeepAlive"],
    "acepto_datos": ["#aceptoNoMolestarCheck"],
    "marco_no_molestar": ["#frameNoMolestar"],
    "boton_nueva_solicitud": ["a.btn-primary:has-text('Ingresar nuevo')"],
    "radio_telefono": ["#customRadioInline1"],
    "radio_correo": ["#customRadioInline2"],
    "tipo_telefono": ["select#tipo-telefono"],
    "campo_telefono": ["input#telefono-sub"],
    "campo_correo_usuario": ["input#mail-pre"],
    "campo_correo_dominio": ["input#mail-sub"],
    "boton_agregar_contacto": ["#ingresar-bloqueo"],
    "lista_telefonos": ["#lista-telefonos-bloqueos li.telefonos", "li.telefonos", ".telefonos"],  # respaldos: platform-9 y WALK.md
    "lista_correos": ["#lista-mails-bloqueos li.correos"],
    "error_contacto": ["#mensajes-error-telefonos > .alert:not(.d-none)"],
    "select_empresa": ["select#empresas"],
    "abrir_buscador_empresa": [
        "select#empresas + .select2 .select2-selection",
        "[aria-labelledby='select2-empresas-container']",
    ],
    "campo_buscar_empresa": [".select2-container--open .select2-search__field"],
    "opciones_empresa": ["#select2-empresas-results li.select2-results__option[role='option']"],
    "boton_agregar_empresa": ["#ingresar-empresa"],
    "lista_empresas": ["#lista-empresas li.empresas"],
    "error_empresa": ["#error-empresas:not(.d-none)"],
    "boton_continuar": ["#fake-continuar"],
    "canales": [
        "input.form-check-input[id*='-empresa_canal_ids_']",
        "input.form-check-input[id$='-empresa]']",
    ],
    "boton_confirmar_telefonos": ["#pills-telefonos button[id^='confirmar-']"],
    "boton_enviar": ["#finalizar"],
    "confirmacion": [
        "#messages-remote-block .alert-success",
        ".alert-success",
        "text=/solicitud ha sido ingresada/i",
    ],
}


def _fuente() -> str:
    return Path(nm.__file__).read_text(encoding="utf-8")


def test_config_selectores_coinciden_con_la_tabla_verificada():
    assert nm.SELECTORES == SELECTORES_ESPERADOS


def test_config_sin_verificar_marca_solo_cu_error_linea_a_linea():
    fuente = _fuente()
    inicio = fuente.index("SELECTORES: dict[str, list[str]] = {")
    fin = fuente.index("\n}\n", inicio)
    clave = None
    marcadas: set[str] = set()
    sin_marca: list[str] = []
    for linea in fuente[inicio:fin].splitlines()[1:]:
        m = re.match(r'\s*"(\w+)":', linea)
        if m:
            clave = m.group(1)
        if "# SIN VERIFICAR" in linea:
            marcadas.add(clave)
        elif clave == "cu_error" and ('"' in linea or "'" in linea):
            sin_marca.append(linea)
    assert marcadas == {"cu_error"}
    assert sin_marca == []


def test_config_bloque_unico_al_inicio_con_todas_las_constantes():
    fuente = _fuente()
    ini = fuente.index("# ====================== CONFIGURACIÓN")
    fin = fuente.index("# ====", ini + 10)
    bloque = fuente[ini:fin]
    for nombre in (
        "PORTAL_URL =", "URL_NO_MOLESTAR =", "SELECTORES_VERIFICADOS =", "FORMATO_RUN =", "SELECTORES:",
        "REGEX_NUMERO_SOLICITUD =", "COLUMNAS:", "TIEMPOS:", "SMTP_HOST, SMTP_PUERTO =", "HOJA =",
    ):
        assert nombre in bloque, nombre
    assert "\ndef " not in fuente[:ini] and "\nclass " not in fuente[:ini]


def test_config_tiempos_tiene_los_valores_de_la_spec():
    assert nm.TIEMPOS == {
        "pausa_entre_empresas": (8, 20),
        "pausa_entre_correos": (15, 40),
        "pausa_entre_acciones": (0.4, 1.2),
        "pausa_entre_busquedas": (1.0, 2.0),
        "timeout_elemento_ms": 15000,
        "timeout_login_s": 300,
        "timeout_confirmacion_ms": 30000,
        "retardo_tecleo_ms": 70,
        "reintentos": 3,
        "backoff_base_s": 4,
    }


def test_config_columnas_tiene_los_alias_de_la_spec():
    assert nm.COLUMNAS == {
        "numero": ["#", "N°", "Nro"],
        "empresa": ["Marca", "Empresa", "Nombre empresa"],
        "razon_social": ["Razón social"],
        "rut_empresa": ["RUT", "RUT empresa"],
        "correo_empresa": ["Correo datos personales", "Correo empresa", "Email empresa"],
        "via": ["Vía de envío", "Via de envio"],
        "formulario_web": ["Formulario web"],
        "canal_alternativo": ["Canal alternativo"],
        "fecha_correo": ["Enviado (fecha)", "Fecha correo"],
        "fecha_sernac": ["No Molestar ingresado", "Fecha SERNAC"],
        "numero_sernac": ["N° solicitud SERNAC", "Numero solicitud SERNAC"],
    }


def test_config_constantes_generales():
    assert nm.PORTAL_URL == "https://www.sernac.cl/app/consumidor/"
    assert nm.URL_NO_MOLESTAR == "https://www.sernac.cl/app/consumidor/index.php?c=home&a=nomolestar"
    assert nm.SELECTORES_VERIFICADOS is True
    assert nm.FORMATO_RUN == "solo_digitos"
    assert (nm.SMTP_HOST, nm.SMTP_PUERTO) == ("smtp.gmail.com", 465)
    assert nm.HOJA == "Empresas"
    assert re.compile(nm.REGEX_NUMERO_SOLICITUD, re.I).groups >= 1


def test_config_selectores_captcha_no_cuentan_el_badge_invisible():
    assert nm.SELECTORES_CAPTCHA
    assert not any("grecaptcha-badge" in s for s in nm.SELECTORES_CAPTCHA)
    assert not set(nm.SELECTORES_CAPTCHA) & set(SELECTORES_QUE_VEN_EL_BADGE)
    anchors = [s for s in nm.SELECTORES_CAPTCHA if "anchor" in s]
    assert anchors and all("size=invisible" in s and ":not(" in s for s in anchors)


def test_importar_el_modulo_no_importa_playwright():
    codigo = "import sys, no_molestar; print('playwright' in sys.modules)"
    r = subprocess.run([sys.executable, "-c", codigo], cwd=RAIZ, capture_output=True, text=True, timeout=60)
    assert r.returncode == 0, r.stderr
    assert r.stdout.strip() == "False"


# =============================================================================
# Normalización y validación
# =============================================================================
@pytest.mark.parametrize(
    "entrada, esperado",
    [("  Vía de   Envío ", "via de envio"), ("Razón Social S.A.", "razon social s a"), ("ÁÉÍÓÚ ñ Ü", "aeiou n u"), (None, ""), (123, "123")],
)
def test_normalizar_quita_tildes_mayusculas_y_puntuacion(entrada, esperado):
    assert nm.normalizar(entrada) == esperado


@pytest.mark.parametrize(
    "variantes",
    [
        ["N° solicitud SERNAC", "Nº solicitud SERNAC", "No. solicitud sernac", "n° SOLICITUD  sernac"],
        ["Vía de envío", "via de envio", "VIA DE ENVÍO", "Vía  de-envío"],
        ["Razón social", "razon social", "RAZÓN SOCIAL"],
        ["Enviado (fecha)", "enviado fecha", "ENVIADO  (FECHA)"],
    ],
)
def test_normalizar_encabezado_hace_equivalentes_las_variantes(variantes):
    assert len({nm.normalizar_encabezado(v) for v in variantes}) == 1


def test_normalizar_encabezado_distingue_columnas_distintas():
    assert nm.normalizar_encabezado("Correo datos personales") != nm.normalizar_encabezado("Correo")
    assert nm.normalizar_encabezado("No Molestar ingresado") != nm.normalizar_encabezado("Enviado (fecha)")
    assert nm.normalizar_encabezado("RUT") != nm.normalizar_encabezado("RUT empresa")


def test_normalizar_encabezado_de_numeral_no_queda_vacio():
    assert nm.normalizar_encabezado("#") != ""
    assert nm.normalizar_encabezado(" # ") == nm.normalizar_encabezado("#")
    assert nm.normalizar_encabezado(None) == "" and nm.normalizar_encabezado("") == ""


@pytest.mark.parametrize("valor", ["Correo", " correo ", "CORREO", "Email", "E-mail", "e mail", "mail"])
def test_es_via_correo_acepta_las_variantes_de_correo(valor):
    assert nm.es_via_correo(valor) is True


@pytest.mark.parametrize("valor", ["Formulario web", "Canal alternativo", "", None, "SERNAC", "Correo y formulario"])
def test_es_via_correo_rechaza_las_demas_vias(valor):
    assert nm.es_via_correo(valor) is False


@pytest.mark.parametrize(
    "rut", ["12.345.678-5", "12345678-5", "123456785", "10.000.013-K", "10.000.013-k", "10000013k", "11.111.111-1", "7.654.321-6"]
)
def test_validar_rut_acepta_ruts_validos(rut):
    assert nm.validar_rut(rut) is True


@pytest.mark.parametrize("rut", ["12.345.678-9", "10.000.013-1", "", "abcdefgh-5", "12.345.678-X", "12.345.678"])
def test_validar_rut_rechaza_ruts_invalidos(rut):
    assert nm.validar_rut(rut) is False


@pytest.mark.parametrize(
    "rut, estilo, esperado",
    [
        ("12345678-5", "puntos", "12.345.678-5"),
        ("123456785", "sin_puntos", "12345678-5"),
        ("12.345.678-5", "solo_digitos", "123456785"),
        ("10.000.013-k", "puntos", "10.000.013-K"),
        ("7654321-6", "puntos", "7.654.321-6"),
    ],
)
def test_formatear_rut_en_cada_estilo(rut, estilo, esperado):
    assert nm.formatear_rut(rut, estilo) == esperado


def test_formatear_rut_estilo_desconocido_lanza_value_error():
    with pytest.raises(ValueError):
        nm.formatear_rut("12345678-5", "con_guiones_raros")


def test_persona_expone_rut_telefono_y_partes_del_correo(persona):
    assert persona.rut_formateado == "12.345.678-5"
    assert persona.telefono_e164 == "+56912345678"
    assert persona.telefono_legible == "+56 9 1234 5678"
    assert persona.correo_usuario == "ana.perez"
    assert persona.correo_dominio == "example.cl"


@pytest.mark.parametrize(
    "texto, esperado",
    [
        ("+56 9 1234 5678", "912345678"),
        ("56912345678", "912345678"),
        ("912345678", "912345678"),
        ("9 1234 5678", "912345678"),
        ("(+56) 9-1234-5678", "912345678"),
        ("+56912345678", "912345678"),
        ("+56 2 2345 6789", "223456789"),
    ],
)
def test_normalizar_telefono_devuelve_nueve_digitos_nacionales(texto, esperado):
    assert nm.normalizar_telefono(texto) == esperado


@pytest.mark.parametrize("texto", ["12345", "abc", "+56 9 1234 56789", "", "+1 555 123 4567"])
def test_normalizar_telefono_invalido_lanza_value_error(texto):
    with pytest.raises(ValueError):
        nm.normalizar_telefono(texto)


@pytest.mark.parametrize("correo", ["ana@example.cl", "ana.perez+spam@sub.example.com", "protecciondedatos@entel.cl"])
def test_validar_correo_acepta_direcciones_basicas(correo):
    assert nm.validar_correo(correo) is True


@pytest.mark.parametrize(
    "correo",
    ["ana@", "ana example.cl", "", "sin-arroba.cl", "ana@example", "ana@example.cl\n", "ana@example.cl\r\nBcc: x@y.cl"],
)
def test_validar_correo_rechaza_invalidos_y_saltos_de_linea(correo):
    assert nm.validar_correo(correo) is False


# =============================================================================
# Directorio SERNAC: candidatos, término de búsqueda y ranking (datos reales de WALK.md)
# =============================================================================
def test_candidatos_busqueda_separa_marca_y_parentesis():
    candidatos = nm.candidatos_busqueda(fila_real(3))
    normal = [_norm(c) for c in candidatos]
    assert normal[0] == "movistar chile"
    assert "telefonica chile" in normal
    assert not any("(" in c or ")" in c for c in candidatos)


def test_candidatos_busqueda_separa_las_partes_con_barra():
    candidatos = nm.candidatos_busqueda(fila_de("Mercado Libre / Mercado Pago Chile", numero=10, via="Canal alternativo"))
    normal = [_norm(c) for c in candidatos]
    assert normal[0] == "mercado libre"
    assert "mercado pago chile" in normal
    assert not any("/" in c for c in candidatos)


def test_candidatos_busqueda_agrega_las_partes_de_la_razon_social():
    normal = [_norm(c) for c in nm.candidatos_busqueda(fila_real(4))]
    assert normal[0] == "entel"
    assert "empresa nacional de telecomunicaciones s a" in normal
    assert "entel pcs telecomunicaciones s a" in normal


def test_candidatos_busqueda_razon_social_con_punto_y_coma():
    fila = fila_de("Zurich Chile", numero=35, razon_social="Zurich Chile Seguros de Vida S.A.; Zurich Chile Seguros Generales S.A.")
    normal = [_norm(c) for c in nm.candidatos_busqueda(fila)]
    assert "zurich chile seguros de vida s a" in normal
    assert "zurich chile seguros generales s a" in normal


def test_candidatos_busqueda_sin_duplicados_ni_vacios():
    candidatos = nm.candidatos_busqueda(fila_real(11))
    assert [_norm(c) for c in candidatos] == ["banco de chile"]
    assert all(c.strip() for c in candidatos)


@pytest.mark.parametrize(
    "texto, esperado",
    [
        (TEXTO_POR_ID[178], "CLARO MOVIL"),
        ("BANCO DE CHILE - BANCO EDWARDS - CITI", "BANCO DE CHILE"),
        ("MERCADO LIBRE- PORTALINMOBILIARIO.COM - MERCADO PAGO - MERCADO PAGO EMISORA", "MERCADO LIBRE- PORTALINMOBILIARIO.COM"),
        (
            "CAT/TARJETAS JUMBO, PARIS, EASY, T.BUS, CENCOSUD VISA, MASTERCARD, BANCO PARIS TARJ. Y CRED.CONSUMO",
            "CAT/TARJETAS JUMBO, PARIS, EASY, T.BUS,",
        ),
        ("SOLVENTA TARJETAS CREDITO FARMACIAS CRUZ VERDE - CHILQUINTA - PREUNIC ANTIGUA- VENTAS IPS EX INP)", "SOLVENTA TARJETAS CREDITO FARMACIAS CRUZ"),
        ("TIENDAS FALABELLA", "TIENDAS FALABELLA"),
        (TEXTO_POR_ID[184], "ENTELPHONE"),
        ("LG (ELECTRONICA)", "LG (ELECTRONICA)"),
    ],
)
def test_termino_busqueda_ejemplos_de_la_spec(texto, esperado):
    assert nm.termino_busqueda(texto) == esperado


def _textos(clave: str) -> list[str]:
    return [t for _, t in RESULTADOS_WALK[clave]]


def _orden(fila, opciones):
    resultado = nm.puntuar_opciones(fila, opciones)
    return resultado, [opciones[i] for i, _ in resultado]


def _sugeridas(resultado, opciones) -> set[str]:
    return {opciones[i] for i, p in resultado if p >= 3}


def test_puntuar_opciones_entel_pone_las_entradas_reales_sobre_el_ruido():
    opciones = _textos("entel")
    resultado, orden = _orden(fila_real(4), opciones)
    buenas = {TEXTO_POR_ID[183], TEXTO_POR_ID[184], TEXTO_POR_ID[432]}
    ruido = {"RESTAURANT HUENTELAUQUEN", "HUENTELAUQUEN", "ZENTELLA MOTOS"}
    assert set(orden[:3]) == buenas
    assert max(orden.index(b) for b in buenas) < min(orden.index(r) for r in ruido)
    assert orden[0] == TEXTO_POR_ID[183]
    assert TEXTO_POR_ID[183] in _sugeridas(resultado, opciones)


def test_puntuar_opciones_claro_sugiere_movil_y_fijo_pero_no_la_clinica():
    opciones = _textos("claro")
    resultado, _ = _orden(fila_real(1), opciones)
    sugeridas = _sugeridas(resultado, opciones)
    assert {TEXTO_POR_ID[178], TEXTO_POR_ID[146420]} <= sugeridas
    assert "CLINICA MIGUEL CLARO" not in sugeridas


def test_puntuar_opciones_banco_falabella_sugiere_solo_las_dos_del_banco():
    opciones = _textos("falabella")
    fila = fila_de("Banco Falabella", numero=19, correo_empresa="datospersonalesbfcl@bancofalabella.cl")
    resultado, orden = _orden(fila, opciones)
    esperadas = {TEXTO_POR_ID[159], TEXTO_POR_ID[146573]}
    assert set(orden[:2]) == esperadas
    assert _sugeridas(resultado, opciones) == esperadas


def test_puntuar_opciones_banco_de_chile_no_sugiere_afubach():
    opciones = _textos("banco de chile")
    resultado, orden = _orden(fila_real(11), opciones)
    assert orden[0] == "BANCO DE CHILE - BANCO EDWARDS - CITI"
    assert _sugeridas(resultado, opciones) == {"BANCO DE CHILE - BANCO EDWARDS - CITI"}


def test_puntuar_opciones_penaliza_la_marca_dentro_de_otra_palabra():
    resultado = dict(nm.puntuar_opciones(fila_de("Entel", numero=4), ["HUENTELAUQUEN", "SUR ENTEL"]))
    assert resultado[0] < 0
    assert resultado[1] > resultado[0]


def test_puntuar_opciones_devuelve_todas_ordenadas_desc_y_empates_por_indice():
    opciones = _textos("claro")
    resultado = nm.puntuar_opciones(fila_real(1), opciones)
    assert sorted(i for i, _ in resultado) == list(range(len(opciones)))
    puntajes = [p for _, p in resultado]
    assert puntajes == sorted(puntajes, reverse=True)
    for (i1, p1), (i2, p2) in zip(resultado, resultado[1:]):
        if p1 == p2:
            assert i1 < i2


# =============================================================================
# Números de solicitud y nombres de archivo
# =============================================================================
@pytest.mark.parametrize(
    "texto, esperado",
    [
        ("N° de solicitud: 1234567", "1234567"),
        ("Nº Solicitud 1234567", "1234567"),
        ("Número de solicitud: AB-12345", "AB-12345"),
        ("Folio: 98765", "98765"),
        ("N° de folio 123456", "123456"),
        ("Código de solicitud: X1Y2Z3", "X1Y2Z3"),
        ("Solicitud N° 123456", "123456"),
        ("FOLIO: 98765", "98765"),
        ("Gracias.\nSu solicitud fue recibida. N° de solicitud: 1234567. Plazo: 7 días.", "1234567"),
    ],
)
def test_extraer_numero_solicitud_variantes(texto, esperado):
    assert nm.extraer_numero_solicitud(texto) == esperado


@pytest.mark.parametrize("texto", ["Su solicitud fue ingresada correctamente", "Folio: 12", ""])
def test_extraer_numero_solicitud_sin_numero_devuelve_none(texto):
    assert nm.extraer_numero_solicitud(texto) is None


@pytest.mark.parametrize("texto", ['Falabella S.A. / Retail: <Chile>?*|"', "..\\..\\Windows\\system32", "../../etc/passwd", "a\tb\nc"])
def test_slug_no_deja_separadores_ni_caracteres_reservados(texto):
    s = nm.slug(texto)
    assert re.fullmatch(r"[a-z0-9-]+", s), s
    assert ".." not in s


def test_slug_translitera_tildes():
    assert nm.slug("Telefónica Ñuñoa") == "telefonica-nunoa"


def test_slug_respeta_el_largo_maximo():
    assert len(nm.slug("x" * 300)) <= 60
    assert len(nm.slug("a b " * 100, largo=20)) <= 20


@pytest.mark.parametrize("texto", ["", "???", "   "])
def test_slug_nunca_queda_vacio(texto):
    assert re.fullmatch(r"[a-z0-9-]+", nm.slug(texto))


@pytest.mark.parametrize("texto", ["CON", "nul", "Com1", "LPT9", "aux", "PRN"])
def test_slug_evita_nombres_reservados_de_windows(texto):
    reservados = {"con", "prn", "aux", "nul"} | {f"com{i}" for i in range(1, 10)} | {f"lpt{i}" for i in range(1, 10)}
    assert nm.slug(texto) not in reservados


# =============================================================================
# Reintentos y pausas
# =============================================================================
class Transitorio(Exception):
    pass


class Permanente(Exception):
    pass


def fallar_n_veces(n: int, exc_tipo=Transitorio, resultado="ok"):
    estado = {"llamadas": 0}

    def fn():
        estado["llamadas"] += 1
        if estado["llamadas"] <= n:
            raise exc_tipo(f"fallo {estado['llamadas']}")
        return resultado

    return fn, estado


def es_transitorio_prueba(exc: BaseException) -> bool:
    return isinstance(exc, Transitorio)


def test_con_reintentos_tiene_exito_tras_errores_transitorios(sleep_falso):
    fn, estado = fallar_n_veces(2)
    r = nm.con_reintentos(fn, intentos=3, base=1, es_transitorio=es_transitorio_prueba, sleep=sleep_falso)
    assert r == "ok"
    assert estado["llamadas"] == 3
    assert len(sleep_falso.llamadas) == 2


def test_con_reintentos_no_reintenta_errores_permanentes(sleep_falso):
    fn, estado = fallar_n_veces(5, exc_tipo=Permanente)
    with pytest.raises(Permanente):
        nm.con_reintentos(fn, intentos=3, base=1, es_transitorio=es_transitorio_prueba, sleep=sleep_falso)
    assert estado["llamadas"] == 1
    assert sleep_falso.llamadas == []


def test_con_reintentos_agotados_relanza_el_ultimo_error(sleep_falso):
    fn, estado = fallar_n_veces(10)
    with pytest.raises(Transitorio, match="fallo 3"):
        nm.con_reintentos(fn, intentos=3, base=1, es_transitorio=es_transitorio_prueba, sleep=sleep_falso)
    assert estado["llamadas"] == 3
    assert len(sleep_falso.llamadas) == 2


def test_con_reintentos_backoff_crece_y_parte_en_la_base(sleep_falso):
    fn, _ = fallar_n_veces(10)
    with pytest.raises(Transitorio):
        nm.con_reintentos(fn, intentos=4, base=2, es_transitorio=es_transitorio_prueba, sleep=sleep_falso)
    esperas = sleep_falso.llamadas
    assert len(esperas) == 3
    assert esperas[0] >= 2
    assert all(b > a for a, b in zip(esperas, esperas[1:]))


def test_con_reintentos_llama_al_reintentar_con_numero_y_excepcion(sleep_falso):
    fn, _ = fallar_n_veces(2)
    vistos = []
    nm.con_reintentos(
        fn, intentos=3, base=1, es_transitorio=es_transitorio_prueba, sleep=sleep_falso,
        al_reintentar=lambda n, exc: vistos.append((n, str(exc))),
    )
    assert vistos == [(1, "fallo 1"), (2, "fallo 2")]


def test_con_reintentos_con_un_intento_no_reintenta(sleep_falso):
    fn, estado = fallar_n_veces(1)
    with pytest.raises(Transitorio):
        nm.con_reintentos(fn, intentos=1, base=1, es_transitorio=es_transitorio_prueba, sleep=sleep_falso)
    assert estado["llamadas"] == 1
    assert sleep_falso.llamadas == []


@pytest.mark.parametrize(
    "exc, esperado",
    [
        (PWTimeoutError("lento"), True),
        (PWError("se cerró el contexto"), True),
        (nm.AbortoUsuario(), False),
        (nm.EmpresaOmitida(), False),
        (nm.ErrorLogin("clave incorrecta"), False),
        (ValueError("otro"), False),
    ],
)
def test_es_transitorio_playwright(exc, esperado):
    assert nm.es_transitorio_playwright(exc) is esperado


@pytest.mark.parametrize(
    "exc, esperado",
    [
        (smtplib.SMTPServerDisconnected("cortado"), True),
        (smtplib.SMTPConnectError(421, b"ocupado"), True),
        (socket.timeout("timeout"), True),
        (TimeoutError("timeout"), True),
        (ConnectionResetError("reset"), True),
        (smtplib.SMTPResponseException(421, b"temporal"), True),
        (smtplib.SMTPDataError(451, b"intente mas tarde"), True),
        (smtplib.SMTPDataError(550, b"rechazado"), False),
        (smtplib.SMTPSenderRefused(553, b"remitente rechazado", "ana@example.cl"), False),
        (smtplib.SMTPRecipientsRefused({"x@y.cl": (550, b"no existe")}), False),
        (smtplib.SMTPAuthenticationError(535, b"5.7.8 Username and Password not accepted"), False),
        (ValueError("otro"), False),
    ],
)
def test_es_transitorio_smtp(exc, esperado):
    assert nm.es_transitorio_smtp(exc) is esperado


def test_pausa_cortes_duerme_una_vez_dentro_del_rango(sleep_falso):
    segundos = nm.pausa_cortes((8, 20), sleep_falso)
    assert len(sleep_falso.llamadas) == 1
    assert 8 <= sleep_falso.llamadas[0] <= 20
    assert segundos == sleep_falso.llamadas[0]


def test_pausa_cortes_con_rango_fijo(sleep_falso):
    assert nm.pausa_cortes((2, 2), sleep_falso) == 2
    assert sleep_falso.llamadas == [2]


def test_pausar_muestra_motivo_con_campana_y_espera_enter(capsys):
    preguntar = PreguntarGuionado([""])
    nm.pausar("Resuelve el CAPTCHA en el navegador", preguntar)
    salida = capsys.readouterr()
    texto = salida.out + salida.err + preguntar.texto
    assert "Resuelve el CAPTCHA en el navegador" in texto
    assert "\a" in texto
    assert preguntar.respuestas == []


@pytest.mark.parametrize("respuesta", ["abortar", "  ABORTAR  ", EOFError, KeyboardInterrupt])
def test_pausar_aborta_con_abortar_eof_o_ctrl_c(respuesta):
    with pytest.raises(nm.AbortoUsuario):
        nm.pausar("Motivo", PreguntarGuionado([respuesta]))


# =============================================================================
# Secretos y logging
# =============================================================================
def test_secreto_oculta_el_valor_en_repr_str_y_format():
    s = nm.Secreto(SECRETO_CU)
    assert repr(s) == "Secreto('***')"
    assert str(s) == "Secreto('***')"
    assert f"{s}" == "Secreto('***')"
    assert SECRETO_CU not in repr([s])
    assert SECRETO_CU not in "%s %r" % (s, s)


def test_secreto_revela_el_valor_solo_con_revelar():
    assert nm.Secreto(SECRETO_CU).revelar() == SECRETO_CU


def test_secreto_bool_indica_si_hay_valor():
    assert nm.Secreto("x")
    assert not nm.Secreto("")


def test_secreto_no_compara_por_valor():
    assert nm.Secreto("abc") != "abc"
    assert nm.Secreto("abc") != nm.Secreto("abc")


def registro(msg, args=(), exc_info=None) -> logging.LogRecord:
    return logging.LogRecord(nm.NOMBRE_LOGGER, logging.INFO, __file__, 1, msg, args, exc_info)


def test_filtro_secretos_censura_el_mensaje():
    filtro = nm.FiltroSecretos([nm.Secreto(SECRETO_CU)])
    r = registro(f"login con {SECRETO_CU} falló")
    assert filtro.filter(r) is True
    assert SECRETO_CU not in r.getMessage()
    assert "***" in r.getMessage()


def test_filtro_secretos_censura_los_argumentos():
    filtro = nm.FiltroSecretos([nm.Secreto(SECRETO_CU)])
    r = registro("clave=%s usuario=%s", (SECRETO_CU, "ana"))
    filtro.filter(r)
    mensaje = r.getMessage()
    assert SECRETO_CU not in mensaje
    assert "ana" in mensaje


def test_filtro_secretos_censura_la_variante_sin_espacios():
    filtro = nm.FiltroSecretos([nm.Secreto(SECRETO_GMAIL)])
    r = registro(f"usando {SECRETO_GMAIL_SIN_ESPACIOS} y también '{SECRETO_GMAIL}'")
    filtro.filter(r)
    assert SECRETO_GMAIL_SIN_ESPACIOS not in r.getMessage()
    assert SECRETO_GMAIL not in r.getMessage()


def test_filtro_secretos_censura_el_texto_de_la_excepcion():
    filtro = nm.FiltroSecretos([nm.Secreto(SECRETO_CU)])
    try:
        raise RuntimeError(f"falló el login con {SECRETO_CU}")
    except RuntimeError:
        info = sys.exc_info()
    r = registro("error inesperado", exc_info=info)
    filtro.filter(r)
    texto = logging.Formatter("%(message)s").format(r)
    assert "RuntimeError" in texto
    assert SECRETO_CU not in texto


def test_filtro_secretos_agregar_censura_secretos_posteriores():
    filtro = nm.FiltroSecretos()
    filtro.agregar(nm.Secreto(SECRETO_CU))
    r = registro(f"x {SECRETO_CU}")
    filtro.filter(r)
    assert SECRETO_CU not in r.getMessage()


@pytest.mark.parametrize("separador", ["?", "&"])
def test_filtro_secretos_censura_el_token_p_del_portal(separador):
    filtro = nm.FiltroSecretos()
    r = registro("marco: %s", (f"https://www.sernac.cl/no-molestar/solicitudes{separador}p={JWT_PORTAL} listo",))
    filtro.filter(r)
    mensaje = r.getMessage()
    assert JWT_PORTAL not in mensaje
    assert f"{separador}p=***" in mensaje


def _vaciar_handlers():
    for lg in (logging.getLogger(nm.NOMBRE_LOGGER), logging.getLogger()):
        for h in lg.handlers:
            h.flush()


def test_configurar_logging_crea_archivo_y_censura_en_archivo_y_consola(tmp_path, capsys):
    directorio = tmp_path / "logs"
    ruta, filtro = nm.configurar_logging(directorio, [nm.Secreto(SECRETO_CU)], verbose=False)
    log = logging.getLogger(nm.NOMBRE_LOGGER)
    log.warning("intento con %s y %s", SECRETO_CU, "otra cosa")
    log.warning(f"texto directo {SECRETO_CU} en ?p={JWT_PORTAL}")
    _vaciar_handlers()
    ruta = Path(ruta)
    contenido = ruta.read_text(encoding="utf-8")
    consola = capsys.readouterr()
    assert ruta.parent.resolve() == directorio.resolve()
    assert re.fullmatch(r"no_molestar_\d{8}-\d{6}\.log", ruta.name)
    assert "intento con" in contenido and "otra cosa" in contenido
    assert "***" in contenido
    for prohibido in (SECRETO_CU, JWT_PORTAL):
        assert prohibido not in contenido
        assert prohibido not in consola.out + consola.err
    assert isinstance(filtro, nm.FiltroSecretos)


def test_configurar_logging_permite_agregar_secretos_despues(tmp_path, capsys):
    ruta, filtro = nm.configurar_logging(tmp_path, [], verbose=True)
    filtro.agregar(nm.Secreto(SECRETO_GMAIL))
    log = logging.getLogger(nm.NOMBRE_LOGGER)
    log.warning("clave: %s", SECRETO_GMAIL_SIN_ESPACIOS)
    log.warning(f"clave: {SECRETO_GMAIL}")
    _vaciar_handlers()
    contenido = Path(ruta).read_text(encoding="utf-8")
    consola = capsys.readouterr()
    for secreto in (SECRETO_GMAIL, SECRETO_GMAIL_SIN_ESPACIOS):
        assert secreto not in contenido
        assert secreto not in consola.out + consola.err


# =============================================================================
# Resumen
# =============================================================================
def test_escribir_resumen_csv_punto_y_coma_utf8_sig_y_neutraliza_formulas(tmp_path):
    resultados = [
        nm.Resultado(5, "Entel", "sernac", "ok", numero="1234567", captura="capturas/x.png"),
        nm.Resultado(2, '=HYPERLINK("http://malo")', "correo", "fallida", detalle="+cmd|' /C calc'!A0"),
        nm.Resultado(3, "Clínica Ñuñoa; Ltda", "correo", "fallida", detalle="@SUMA(1+1)"),
        nm.Resultado(4, "-2+3", "sernac", "incierta"),
    ]
    ruta = nm.escribir_resumen_csv(resultados, tmp_path / "resumen.csv")
    datos = Path(ruta).read_bytes()
    assert datos.startswith(BOM)
    texto = datos.decode("utf-8-sig")
    filas = list(csv.reader(io.StringIO(texto), delimiter=";"))
    encabezado = [c.lower() for c in filas[0]]
    assert len(encabezado) >= 7
    assert {"empresa", "fase", "estado"} <= set(encabezado)
    assert all(len(f) == len(filas[0]) for f in filas[1:])
    celdas_csv = [c for fila in filas[1:] for c in fila]
    assert "Entel" in celdas_csv and "1234567" in celdas_csv and "Clínica Ñuñoa; Ltda" in celdas_csv
    assert "'=HYPERLINK(\"http://malo\")" in celdas_csv
    assert "'-2+3" in celdas_csv
    assert not [c for c in celdas_csv if c and c[0] in "=+-@"]


def test_resumen_texto_conteos_fallas_accion_manual_y_sin_entrada():
    filas = [fila_real(1), fila_real(3), fila_real(4), fila_real(5), fila_real(18)]
    resultados = [
        nm.Resultado(2, "Claro Chile", "sernac", "ok", numero="1234567"),
        nm.Resultado(5, "Entel", "sernac", "fallida", detalle="Timeout en el portal"),
        nm.Resultado(4, "Movistar Chile (Telefónica Chile)", "sernac", "incierta", detalle="confirmación no detectada"),
        nm.Resultado(19, "Banco Security", "sernac", "sin_coincidencia", detalle="sin entrada en SERNAC"),
        nm.Resultado(6, "WOM Chile", "sernac", "ya_hecha"),
        nm.Resultado(2, "Claro Chile", "correo", "ok"),
    ]
    texto = nm.resumen_texto(resultados, filas)
    bajo = texto.lower()
    assert "sernac" in bajo and "correo" in bajo
    for etiqueta in ("fallidas", "inciertas", "sin coincidencia", "omitidas", "ya hechas"):
        assert etiqueta in bajo
    assert "Timeout en el portal" in texto and "confirmación no detectada" in texto
    assert "acción manual pendiente" in bajo
    assert "https://legalmovistar.mvch.cl/" in texto
    assert "Centro de ayuda WOM" in texto
    assert "sin entrada en sernac" in bajo and "Banco Security" in texto
    assert "https://denuncias.americamovil.com/" not in texto  # Claro va por correo: no es acción manual


# =============================================================================
# CLI
# =============================================================================
def test_parse_args_valores_por_defecto():
    a = nm.parse_args([])
    assert a.dry_run is False
    assert a.mapear is False
    assert a.solo is None
    assert a.limite is None
    assert not a.empresa
    assert a.forzar is False
    assert a.explorar is False
    assert a.verbose is False
    assert a.confirmar is (not nm.SELECTORES_VERIFICADOS)
    assert Path(a.tracker) == RAIZ / "tracker_no_molestar_50_empresas.xlsx"
    assert Path(a.config) == RAIZ / "config.toml"
    assert Path(a.plantilla) == RAIZ / "plantilla.txt"
    assert Path(a.mapeo) == RAIZ / "mapeo_sernac.csv"
    assert Path(a.salida) == RAIZ / "salida"


def test_parse_args_opciones():
    a = nm.parse_args(
        ["--dry-run", "--solo", "correo", "--limite", "3", "--empresa", "entel", "--empresa", "banco de chile",
         "--forzar", "--sin-confirmar", "--salida", "otra", "--mapeo", "m.csv", "-v"]
    )
    assert a.dry_run is True
    assert a.solo == "correo"
    assert a.limite == 3
    assert list(a.empresa) == ["entel", "banco de chile"]
    assert a.forzar is True
    assert a.confirmar is False
    assert Path(a.salida) == Path("otra")
    assert Path(a.mapeo) == Path("m.csv")
    assert a.verbose is True


def test_parse_args_mapear_y_confirmar_explicito():
    assert nm.parse_args(["--mapear"]).mapear is True
    assert nm.parse_args(["--confirmar"]).confirmar is True


@pytest.mark.parametrize(
    "argv",
    [
        ["--confirmar", "--sin-confirmar"],
        ["--explorar", "--dry-run"],
        ["--mapear", "--dry-run"],
        ["--mapear", "--explorar"],
        ["--solo", "fax"],
        ["--limite", "x"],
    ],
)
def test_parse_args_rechaza_combinaciones_invalidas(argv, capsys):
    with pytest.raises(SystemExit) as e:
        nm.parse_args(argv)
    assert e.value.code == 2


def test_parse_args_help_lista_las_opciones(capsys):
    with pytest.raises(SystemExit) as e:
        nm.parse_args(["--help"])
    assert e.value.code == 0
    ayuda = capsys.readouterr().out
    for opcion in ("--dry-run", "--mapear", "--solo", "--limite", "--empresa", "--forzar", "--explorar", "--mapeo", "--salida"):
        assert opcion in ayuda


# =============================================================================
# config.toml
# =============================================================================
@pytest.mark.parametrize(
    "texto",
    [
        'name = "Ana Pérez Soto"\nrun = "12345678-5"\nphone = "912345678"\nemail = "ana.perez@example.cl"\n',
        '[persona]\nnombre_completo = "Ana Pérez Soto"\nrut = "123456785"\n"teléfono" = "+56912345678"\nmail = "ana.perez@example.cl"\n',
        '[persona]\nnombre = "Ana Pérez Soto"\nrun = "12.345.678-5"\ncelular = "(+56) 9-1234-5678"\n"e-mail" = "ana.perez@example.cl"\n',
        'nombre = "Ana Pérez Soto"\nrut = "12.345.678-5"\nfono = "9 1234 5678"\ncorreo = "ana.perez@example.cl"\n',
        'nombre = "Ana Pérez Soto"\nrut = "12.345.678-5"\nmovil = "56912345678"\ncorreo = "ana.perez@example.cl"\n',
    ],
)
def test_cargar_config_acepta_alias_en_la_raiz_o_en_persona(escribir, texto):
    cfg = nm.cargar_config(escribir("config.toml", texto))
    assert cfg.persona.nombre == "Ana Pérez Soto"
    assert cfg.persona.rut == "12345678-5"
    assert cfg.persona.telefono == "912345678"
    assert cfg.persona.correo == "ana.perez@example.cl"
    assert cfg.gmail_usuario == "ana.perez@example.cl"


def test_cargar_config_usa_gmail_usuario_si_existe(escribir):
    cfg = nm.cargar_config(escribir("config.toml", CONFIG_VALIDA))
    assert cfg.gmail_usuario == GMAIL_USUARIO
    assert cfg.persona.correo == CORREO_PERSONA


def test_cargar_config_ignora_claves_secretas_con_advertencia_sin_mostrar_valores(escribir, caplog, capsys):
    caplog.set_level(logging.WARNING)
    texto = (
        'token = "tok-raiz-998877"\n\n'
        + CONFIG_VALIDA.replace('correo = "ana.perez@example.cl"\n', f'correo = "ana.perez@example.cl"\nclave = "{SECRETO_CU}"\n')
        + f'app_password = "{SECRETO_GMAIL}"\n'
    )
    cfg = nm.cargar_config(escribir("config.toml", texto))
    salida = capsys.readouterr()
    avisos = caplog.text + salida.out + salida.err
    assert cfg.gmail_usuario == GMAIL_USUARIO
    for clave in ("clave", "app_password", "token"):
        assert clave in avisos
    for valor in (SECRETO_CU, SECRETO_GMAIL, "tok-raiz-998877"):
        assert valor not in avisos
        assert valor not in repr(cfg)


def test_cargar_config_reune_todos_los_errores_de_validacion(escribir):
    texto = '[persona]\nnombre = ""\nrut = "12.345.678-9"\ntelefono = "123"\ncorreo = "no-es-un-correo"\n'
    with pytest.raises(nm.ErrorConfig) as e:
        nm.cargar_config(escribir("config.toml", texto))
    mensaje = str(e.value).lower()
    for campo in ("nombre", "rut", "tel", "correo"):
        assert campo in mensaje


def test_cargar_config_rechaza_el_ejemplo_sin_editar(tmp_path):
    copia = tmp_path / "config.toml"
    shutil.copyfile(Path(nm.__file__).with_name("config.example.toml"), copia)
    with pytest.raises(nm.ErrorConfig, match="ejemplo"):
        nm.cargar_config(copia)


@pytest.mark.parametrize(
    "rut, correo",
    [("11.111.111-1", "ana.perez@gmail.com"), ("12.345.678-5", "Maria.Perez@example.cl")],
)
def test_cargar_config_rechaza_rut_o_correo_del_ejemplo(escribir, rut, correo):
    texto = f'[persona]\nnombre = "Ana"\nrut = "{rut}"\ntelefono = "+56 9 8765 4321"\ncorreo = "{correo}"\n'
    with pytest.raises(nm.ErrorConfig, match="ejemplo"):
        nm.cargar_config(escribir("config.toml", texto))


def test_cargar_config_sin_datos_lanza_error_config(escribir):
    with pytest.raises(nm.ErrorConfig):
        nm.cargar_config(escribir("config.toml", "# vacío\n"))


def test_cargar_config_archivo_inexistente_lanza_error_config(tmp_path):
    with pytest.raises(nm.ErrorConfig):
        nm.cargar_config(tmp_path / "no-existe.toml")


def test_cargar_config_toml_invalido_lanza_error_config(escribir):
    with pytest.raises(nm.ErrorConfig):
        nm.cargar_config(escribir("config.toml", "nombre = \n[[[ roto"))


# =============================================================================
# plantilla.txt, marcadores y contexto
# =============================================================================
def test_cargar_plantilla_con_bom_separa_asunto_y_cuerpo(escribir):
    p = nm.cargar_plantilla(escribir("plantilla.txt", "Asunto:  Hola {empresa}  \n\n\nEstimados de {empresa}:\nTexto.\n", bom=True))
    assert p.asunto == "Hola {empresa}"
    assert p.cuerpo.startswith("Estimados de {empresa}:")
    assert "Texto." in p.cuerpo


@pytest.mark.parametrize("linea", ["Subject: Hola {empresa}", "ASUNTO: Hola {empresa}", "subject:Hola {empresa}"])
def test_cargar_plantilla_acepta_subject_y_mayusculas(escribir, linea):
    p = nm.cargar_plantilla(escribir("plantilla.txt", f"{linea}\nCuerpo para {{empresa}}\n"))
    assert p.asunto == "Hola {empresa}"
    assert p.cuerpo.startswith("Cuerpo para {empresa}")


def test_cargar_plantilla_ignora_lineas_vacias_antes_del_asunto(escribir):
    p = nm.cargar_plantilla(escribir("plantilla.txt", "\n\n  \nAsunto: X\n\nCuerpo\n"))
    assert p.asunto == "X"
    assert p.cuerpo.startswith("Cuerpo")


def test_cargar_plantilla_sin_asunto_usa_el_por_defecto_y_advierte(escribir, caplog, capsys):
    caplog.set_level(logging.WARNING)
    p = nm.cargar_plantilla(escribir("plantilla.txt", "\nEstimados de {empresa}:\nTexto.\n"))
    salida = capsys.readouterr()
    assert p.asunto == nm.ASUNTO_POR_DEFECTO
    assert "{nombre}" in p.asunto
    assert p.cuerpo.startswith("Estimados de {empresa}:")
    assert "asunto" in (caplog.text + salida.out + salida.err).lower()


def test_cargar_plantilla_inexistente_lanza_error_plantilla(tmp_path):
    with pytest.raises(nm.ErrorPlantilla):
        nm.cargar_plantilla(tmp_path / "no-existe.txt")


def test_placeholders_devuelve_nombres_tal_como_estan_escritos():
    texto = "Hola {nombre}, de {Correo Empresa}. {{no}} {fecha} {persona.__class__}"
    assert nm.placeholders(texto) == {"nombre", "Correo Empresa", "fecha", "persona.__class__"}


CONTEXTO = {"nombre": "Ana", "empresa": "Entel", "correo_empresa": "protecciondedatos@entel.cl", "fecha": "22-10-2026"}


def test_renderizar_sustituye_marcadores():
    assert nm.renderizar("Hola {nombre} de {empresa}", CONTEXTO) == "Hola Ana de Entel"


def test_renderizar_busca_claves_normalizadas():
    assert nm.renderizar("{Nombre} <{CORREO EMPRESA}> {correo-empresa}", CONTEXTO) == (
        "Ana <protecciondedatos@entel.cl> protecciondedatos@entel.cl"
    )


def test_renderizar_llaves_dobles_son_literales():
    assert nm.renderizar("{{literal}} {nombre} }}", CONTEXTO) == "{literal} Ana }"


def test_renderizar_marcadores_desconocidos_listan_sus_nombres():
    with pytest.raises(nm.ErrorPlantilla) as e:
        nm.renderizar("Hola {desconocido} y {otro_mas}", CONTEXTO)
    assert "desconocido" in str(e.value)
    assert "otro_mas" in str(e.value)


@pytest.mark.parametrize("texto", ["{persona.__class__}", "{nombre.__class__}", "{0}", "{nombre!r}", "{empresa[0]}"])
def test_renderizar_no_permite_acceso_a_atributos_ni_indices(texto):
    with pytest.raises(nm.ErrorPlantilla):
        nm.renderizar(texto, CONTEXTO)


def test_renderizar_no_reexpande_valores_con_llaves():
    assert nm.renderizar("Empresa: {empresa}", dict(CONTEXTO, empresa="ACME {nombre}")) == "Empresa: ACME {nombre}"


def test_renderizar_no_interpreta_barras_invertidas_del_valor():
    ctx = dict(CONTEXTO, nombre=r"C:\temp\1 \g<0>")
    assert nm.renderizar("Ruta {nombre}", ctx) == r"Ruta C:\temp\1 \g<0>"


def test_contexto_para_fila_real_con_encabezados_empresa_y_persona(persona):
    fila = fila_real(4)
    fila.valores["Notas"] = None
    ctx = nm.contexto_para(fila, persona, HOY)
    assert ctx["rut"] == "12.345.678-5"  # la persona gana a la columna «RUT» de la empresa
    assert ctx["rut_empresa"] == "92.580.000-7 / 96.806.980-2"
    assert ctx["empresa"] == "Entel" and ctx["marca"] == "Entel"
    assert ctx["razon_social"] == "Empresa Nacional de Telecomunicaciones S.A. / Entel PCS Telecomunicaciones S.A."
    assert ctx["correo_empresa"] == "protecciondedatos@entel.cl"
    assert ctx["correo_datos_personales"] == "protecciondedatos@entel.cl"
    assert ctx["correo"] == CORREO_PERSONA
    assert ctx["nombre"] == "Ana Pérez Soto"
    assert ctx["telefono"] == "+56 9 1234 5678"
    assert ctx["fecha"] == "22-10-2026"
    assert ctx["sector"] == "Telecom"
    assert ctx["via_de_envio"] == "Correo"
    assert ctx["notas"] == "" and ctx["estado"] == ""
    assert "" not in ctx
    assert all(isinstance(v, str) for v in ctx.values())


def test_contexto_para_datos_de_empresa_y_persona_ganan_a_encabezados(persona):
    fila = fila_de(
        "Entel", numero=4, razon_social="Entel PCS", correo_empresa="protecciondedatos@entel.cl",
        valores={"#": 4, "Empresa": "Nombre viejo", "Fecha": "01-01-2020", "Nombre": "Otro", "Prioridad": 3},
    )
    ctx = nm.contexto_para(fila, persona, HOY)
    assert ctx["empresa"] == "Entel"
    assert ctx["fecha"] == "22-10-2026"
    assert ctx["nombre"] == "Ana Pérez Soto"
    assert ctx["prioridad"] == "3"


# =============================================================================
# filtrar_pendientes
# =============================================================================
@pytest.fixture
def filas_mixtas():
    return [
        fila_real(1),
        fila_real(3),
        fila_real(4, sernac_hecho=True),
        fila_real(5),
        fila_real(11, correo_hecho=True),
    ]


def test_filtrar_pendientes_sernac_considera_todas_las_empresas(filas_mixtas):
    pendientes, hechas = nm.filtrar_pendientes(filas_mixtas, "sernac")
    assert [f.numero for f in pendientes] == [1, 3, 5, 11]
    assert [(r.numero_fila, r.empresa, r.fase, r.estado) for r in hechas] == [(5, "Entel", "sernac", "ya_hecha")]


def test_filtrar_pendientes_correo_solo_via_correo(filas_mixtas):
    pendientes, hechas = nm.filtrar_pendientes(filas_mixtas, "correo")
    assert [f.numero for f in pendientes] == [1, 4]
    assert [(r.empresa, r.fase, r.estado) for r in hechas] == [("Banco de Chile", "correo", "ya_hecha")]


def test_filtrar_pendientes_forzar_incluye_las_ya_hechas(filas_mixtas):
    pendientes, hechas = nm.filtrar_pendientes(filas_mixtas, "sernac", forzar=True)
    assert [f.numero for f in pendientes] == [1, 3, 4, 5, 11]
    assert hechas == []


def test_filtrar_pendientes_por_empresa_normalizada(filas_mixtas):
    pendientes, _ = nm.filtrar_pendientes(filas_mixtas, "sernac", empresas=["CLARO", "telefónica"])
    assert [f.numero for f in pendientes] == [1, 3]


def test_filtrar_pendientes_limite_solo_acota_pendientes(filas_mixtas):
    pendientes, hechas = nm.filtrar_pendientes(filas_mixtas, "sernac", limite=1)
    assert [f.numero for f in pendientes] == [1]
    assert [r.empresa for r in hechas] == ["Entel"]


# =============================================================================
# Tracker (estructura real)
# =============================================================================
def test_tracker_real_lee_50_empresas_y_excluye_el_resumen(tracker_real):
    filas = nm.Tracker(tracker_real).filas()
    assert len(filas) == 50
    assert [f.numero for f in filas] == list(range(1, 51))
    assert [f.numero_fila for f in filas] == list(range(2, 52))
    assert sum(1 for f in filas if f.es_correo) == 30
    assert not any(_norm(f.empresa) == "resumen" for f in filas)
    assert Counter(f.via for f in filas) == {"Correo": 30, "Formulario web": 2, "Canal alternativo": 18}


def test_tracker_real_campos_de_la_fila_entel(tracker_real):
    entel = next(f for f in nm.Tracker(tracker_real).filas() if f.numero == 4)
    assert entel.numero_fila == 5
    assert entel.empresa == "Entel"
    assert entel.razon_social == "Empresa Nacional de Telecomunicaciones S.A. / Entel PCS Telecomunicaciones S.A."
    assert entel.rut_empresa == "92.580.000-7 / 96.806.980-2"
    assert entel.correo_empresa == "protecciondedatos@entel.cl"
    assert entel.via == "Correo"
    assert entel.formulario_web == ""
    assert entel.canal_alternativo == "Mi Entel (web), App Entel, sucursales Entel, call center"
    assert (entel.sernac_hecho, entel.correo_hecho) == (False, False)
    assert entel.valores["Sector"] == "Telecom"
    assert entel.valores["Estado"] == "" and entel.valores["Plazo respuesta correo"] == ""


def test_tracker_real_fila_sin_correo_tiene_textos_vacios(tracker_real):
    movistar = next(f for f in nm.Tracker(tracker_real).filas() if f.numero == 3)
    assert movistar.correo_empresa == "" and movistar.razon_social == "" and movistar.rut_empresa == ""
    assert movistar.formulario_web == "https://legalmovistar.mvch.cl/"
    assert movistar.es_correo is False


def test_tracker_detecta_encabezados_en_la_fila_3(fabrica_tracker):
    path = fabrica_tracker(filas_reales(1, 4), fila_encabezado=3)
    filas = nm.Tracker(path).filas()
    assert [(f.empresa, f.numero_fila) for f in filas] == [("Claro Chile", 4), ("Entel", 5)]


def test_tracker_acepta_alias_de_encabezados_sin_columna_numero(fabrica_tracker):
    encabezados = ["Empresa", "Via de envio", "Correo empresa", "Fecha correo", "Fecha SERNAC", "Numero solicitud SERNAC"]
    path = fabrica_tracker(
        [{"Empresa": "Entel", "Via de envio": "Correo", "Correo empresa": "protecciondedatos@entel.cl"}],
        encabezados=encabezados, con_resumen=False,
    )
    (fila,) = nm.Tracker(path).filas()
    assert (fila.numero, fila.empresa, fila.via, fila.correo_empresa) == (None, "Entel", "Correo", "protecciondedatos@entel.cl")
    assert (fila.razon_social, fila.formulario_web, fila.canal_alternativo) == ("", "", "")


def test_tracker_se_detiene_en_la_primera_fila_vacia(fabrica_tracker):
    path = fabrica_tracker(filas_reales(1, 3) + [{}] + filas_reales(4), con_resumen=False)
    assert [f.numero for f in nm.Tracker(path).filas()] == [1, 3]


def test_tracker_omite_filas_con_numero_no_entero(fabrica_tracker):
    filas = filas_reales(1, 3, 4)
    filas[1]["#"] = "x"
    filas[2]["#"] = 4.0
    path = fabrica_tracker(filas, con_resumen=False)
    assert [(f.numero, f.empresa) for f in nm.Tracker(path).filas()] == [(1, "Claro Chile"), (4, "Entel")]


def test_tracker_sin_fila_en_blanco_igual_excluye_el_resumen(fabrica_tracker):
    path = fabrica_tracker(filas_reales(1, 3, 4), fila_en_blanco=False)
    filas = nm.Tracker(path).filas()
    assert [f.numero for f in filas] == [1, 3, 4]


def test_tracker_marca_lo_ya_hecho(fabrica_tracker):
    filas = filas_reales(1, 3, 4, 11)
    filas[0]["No Molestar ingresado"] = datetime(2026, 10, 1)
    filas[1]["N° solicitud SERNAC"] = "999999"
    filas[2]["Enviado (fecha)"] = datetime(2026, 10, 2)
    path = fabrica_tracker(filas)
    claro, movistar, entel, chile = nm.Tracker(path).filas()
    assert (claro.sernac_hecho, claro.correo_hecho) == (True, False)
    assert (movistar.sernac_hecho, movistar.correo_hecho) == (True, False)
    assert (entel.sernac_hecho, entel.correo_hecho) == (False, True)
    assert (chile.sernac_hecho, chile.correo_hecho) == (False, False)


def test_tracker_sin_la_hoja_lista_las_hojas_existentes(fabrica_tracker):
    path = fabrica_tracker(filas_reales(1), hoja="Hoja1")
    with pytest.raises(nm.ErrorTracker) as e:
        nm.Tracker(path)
    assert "Hoja1" in str(e.value)


def test_tracker_sin_via_de_envio_lanza_error_que_la_nombra(fabrica_tracker):
    encabezados = [h for h in ENCABEZADOS_REALES if h != "Vía de envío"]
    path = fabrica_tracker([{"Marca": "Entel"}], encabezados=encabezados, con_resumen=False)
    with pytest.raises(nm.ErrorTracker) as e:
        nm.Tracker(path)
    assert "vía de envío" in str(e.value).lower()


@pytest.mark.parametrize("faltante", ["N° solicitud SERNAC", "No Molestar ingresado", "Enviado (fecha)"])
def test_tracker_sin_columna_de_salida_lanza_error_y_no_la_crea(fabrica_tracker, faltante):
    encabezados = [h for h in ENCABEZADOS_REALES if h != faltante]
    path = fabrica_tracker([{"Marca": "Entel", "Vía de envío": "Correo"}], encabezados=encabezados, con_resumen=False)
    antes = path.read_bytes()
    with pytest.raises(nm.ErrorTracker) as e:
        nm.Tracker(path)
    assert faltante in str(e.value)
    assert path.read_bytes() == antes


def test_tracker_sin_fila_de_encabezados_lanza_error(fabrica_tracker):
    path = fabrica_tracker([{"Nombre": "A", "Tipo": "Correo"}], encabezados=["Nombre", "Tipo"], con_resumen=False)
    with pytest.raises(nm.ErrorTracker):
        nm.Tracker(path)


def test_tracker_titulo_con_la_palabra_empresas_no_es_la_fila_de_encabezados(fabrica_tracker):
    path = fabrica_tracker(filas_reales(4), fila_encabezado=2)
    (fila,) = nm.Tracker(path).filas()
    assert (fila.empresa, fila.numero_fila) == ("Entel", 3)


def test_tracker_inexistente_lanza_error_tracker(tmp_path):
    with pytest.raises(nm.ErrorTracker):
        nm.Tracker(tmp_path / "no-existe.xlsx")


def test_tracker_inexistente_indica_copiar_la_plantilla(tmp_path):
    with pytest.raises(nm.ErrorTracker, match=re.escape(nm.TRACKER_PLANTILLA)):
        nm.Tracker(tmp_path / "tracker_no_molestar_50_empresas.xlsx")


def test_tracker_marcar_sernac_escribe_fecha_y_numero_solo_al_guardar(fabrica_tracker):
    path = fabrica_tracker(filas_reales(1, 4))
    antes = path.read_bytes()
    t = nm.Tracker(path)
    t.marcar_sernac(t.filas()[1], HOY, "1234567")
    assert path.read_bytes() == antes
    t.guardar(preguntar=PreguntarGuionado())
    fecha = celda(path, 3, "No Molestar ingresado")
    numero = celda(path, 3, "N° solicitud SERNAC")
    assert es_fecha(fecha.value, HOY)
    assert fecha.number_format == FORMATO_FECHA_REAL
    assert numero.value == "1234567" and numero.data_type == "s"
    assert vacio(valor_celda(path, 3, "Enviado (fecha)"))
    assert vacio(valor_celda(path, 2, "No Molestar ingresado"))


def test_tracker_marcar_sernac_sin_numero_deja_la_celda_del_numero_vacia(fabrica_tracker):
    # El portal real no muestra N° (solo «Su solicitud ha sido ingresada…»): basta la fecha
    path = fabrica_tracker(filas_reales(1))
    t = nm.Tracker(path)
    t.marcar_sernac(t.filas()[0], HOY, None)
    t.guardar(preguntar=PreguntarGuionado())
    assert es_fecha(valor_celda(path, 2, "No Molestar ingresado"), HOY)
    assert vacio(valor_celda(path, 2, "N° solicitud SERNAC"))


def test_tracker_marcar_sernac_sin_numero_no_borra_un_numero_anotado(fabrica_tracker):
    path = fabrica_tracker(filas_reales(1))
    t = nm.Tracker(path)
    t.marcar_sernac(t.filas()[0], HOY, "123456")
    t.marcar_sernac(t.filas()[0], HOY, None)  # p. ej. --forzar: no se pierde lo que ya estaba
    t.guardar(preguntar=PreguntarGuionado())
    assert valor_celda(path, 2, "N° solicitud SERNAC") == "123456"


def test_tracker_fecha_en_celda_general_usa_formato_dd_mm_yyyy(fabrica_tracker):
    path = fabrica_tracker(filas_reales(1), formatos_fecha=False)
    t = nm.Tracker(path)
    t.marcar_correo(t.filas()[0], HOY)
    t.guardar(preguntar=PreguntarGuionado())
    c = celda(path, 2, "Enviado (fecha)")
    assert es_fecha(c.value, HOY)
    assert c.number_format == "dd-mm-yyyy"


def test_tracker_marcar_correo_escribe_solo_enviado_fecha(fabrica_tracker):
    path = fabrica_tracker(filas_reales(1))
    t = nm.Tracker(path)
    t.marcar_correo(t.filas()[0], HOY)
    t.guardar(preguntar=PreguntarGuionado())
    assert es_fecha(valor_celda(path, 2, "Enviado (fecha)"), HOY)
    assert vacio(valor_celda(path, 2, "No Molestar ingresado"))
    assert vacio(valor_celda(path, 2, "N° solicitud SERNAC"))


@pytest.mark.parametrize("columna, accion", [("No Molestar ingresado", "sernac"), ("N° solicitud SERNAC", "sernac"), ("Enviado (fecha)", "correo")])
def test_tracker_se_niega_a_pisar_una_formula(fabrica_tracker, columna, accion):
    filas = filas_reales(1)
    filas[0][columna] = "=TODAY()"
    path = fabrica_tracker(filas)
    t = nm.Tracker(path)
    fila = t.filas()[0]
    with pytest.raises(nm.ErrorTracker):
        if accion == "sernac":
            t.marcar_sernac(fila, HOY, "1234567")
        else:
            t.marcar_correo(fila, HOY)
    assert valor_celda(path, 2, columna) == "=TODAY()"


def test_tracker_real_solo_cambian_las_tres_celdas_de_salida(tracker_real, tmp_path):
    pristino = tmp_path / "pristino.xlsx"
    shutil.copy2(tracker_real, pristino)
    t = nm.Tracker(tracker_real)
    entel = next(f for f in t.filas() if f.numero == 4)
    t.marcar_sernac(entel, HOY, "1234567")
    t.marcar_correo(entel, HOY)
    t.guardar(preguntar=PreguntarGuionado())
    antes, despues = celdas(pristino), celdas(tracker_real)
    cambiadas = {k for k in antes.keys() | despues.keys() if antes.get(k) != despues.get(k)}
    assert cambiadas == {("Empresas", "L5"), ("Empresas", "O5"), ("Empresas", "P5")}
    wb_a, wb_d = openpyxl.load_workbook(pristino), openpyxl.load_workbook(tracker_real)
    assert wb_d.sheetnames == wb_a.sheetnames
    a, d = wb_a["Empresas"], wb_d["Empresas"]
    assert d.max_column == a.max_column == 21
    assert d.freeze_panes == a.freeze_panes
    assert d.auto_filter.ref == a.auto_filter.ref
    assert [str(v.sqref) for v in d.data_validations.dataValidation] == [str(v.sqref) for v in a.data_validations.dataValidation]
    assert sorted(str(r.sqref) for r in d.conditional_formatting) == sorted(str(r.sqref) for r in a.conditional_formatting)
    assert d["O5"].number_format == a["O5"].number_format


def test_tracker_guardar_es_atomico_con_os_replace(monkeypatch, tmp_path):
    carpeta = tmp_path / "solo"
    carpeta.mkdir()
    path = crear_tracker_real(carpeta / "tracker.xlsx", filas_reales(1))
    real = os.replace
    llamadas = []

    def espia(src, dst, *args, **kwargs):
        llamadas.append((Path(src), Path(dst)))
        return real(src, dst, *args, **kwargs)

    monkeypatch.setattr(os, "replace", espia)
    t = nm.Tracker(path)
    t.marcar_sernac(t.filas()[0], HOY, "1")
    devuelto = t.guardar(preguntar=PreguntarGuionado())
    assert any(dst.resolve() == path.resolve() and src.resolve() != path.resolve() for src, dst in llamadas)
    assert Path(devuelto).resolve() == path.resolve()
    assert sorted(p.name for p in carpeta.iterdir()) == [path.name]


def _replace_que_falla(path: Path, veces: int):
    real = os.replace
    estado = {"fallos": 0}

    def falso(src, dst, *args, **kwargs):
        if Path(dst).resolve() == path.resolve() and (veces < 0 or estado["fallos"] < veces):
            estado["fallos"] += 1
            raise PermissionError(13, "El proceso no tiene acceso al archivo porque está siendo utilizado por otro proceso")
        return real(src, dst, *args, **kwargs)

    return falso, estado


def test_tracker_guardar_con_excel_abierto_pide_cerrar_y_reintenta(fabrica_tracker, monkeypatch, capsys):
    path = fabrica_tracker(filas_reales(1))
    falso, estado = _replace_que_falla(path, veces=1)
    monkeypatch.setattr(os, "replace", falso)
    t = nm.Tracker(path)
    t.marcar_correo(t.filas()[0], HOY)
    preguntar = PreguntarGuionado([""])
    t.guardar(preguntar=preguntar)
    salida = capsys.readouterr()
    assert estado["fallos"] == 1
    assert preguntar.respuestas == []
    assert "excel" in (preguntar.texto + salida.out + salida.err).lower()
    assert es_fecha(valor_celda(path, 2, "Enviado (fecha)"), HOY)


def test_tracker_guardar_con_excel_abierto_y_eof_guarda_pendiente(fabrica_tracker, monkeypatch):
    path = fabrica_tracker(filas_reales(1))
    antes = path.read_bytes()
    falso, _ = _replace_que_falla(path, veces=-1)
    monkeypatch.setattr(os, "replace", falso)
    t = nm.Tracker(path)
    t.marcar_correo(t.filas()[0], HOY)
    devuelto = t.guardar(preguntar=PreguntarGuionado([EOFError]))
    pendiente = path.with_name(f"{path.stem}.pendiente.xlsx")
    assert pendiente.exists()
    assert Path(devuelto).resolve() == pendiente.resolve()
    assert path.read_bytes() == antes
    assert es_fecha(valor_celda(pendiente, 2, "Enviado (fecha)"), HOY)


def test_tracker_respaldar_copia_con_marca_de_tiempo(fabrica_tracker, tmp_path):
    path = fabrica_tracker(filas_reales(1))
    copia = Path(nm.Tracker(path).respaldar(tmp_path / "respaldos"))
    assert copia.exists()
    assert copia.parent.resolve() == (tmp_path / "respaldos").resolve()
    assert re.fullmatch(r"tracker_\d{8}-\d{6}\.xlsx", copia.name)
    assert copia.read_bytes() == path.read_bytes()


# =============================================================================
# Mapeo SERNAC (mapeo_sernac.csv)
# =============================================================================
ENCABEZADO_MAPEO = "#;Marca;ID SERNAC;Entrada SERNAC;Nota"


def _filas_csv(path: Path) -> list[dict[str, str]]:
    texto = path.read_text(encoding="utf-8-sig")
    primera = texto.splitlines()[0]
    delimitador = ";" if primera.count(";") >= primera.count(",") else ","
    return list(csv.DictReader(io.StringIO(texto), delimiter=delimitador))


def test_mapeo_real_tiene_50_empresas_y_67_entradas(mapeo_real):
    m = nm.Mapeo.cargar(mapeo_real)
    entradas = {n: m.entradas(fila_de(f"Empresa {n}", numero=n)) for n in range(1, 51)}
    assert all(v is not None for v in entradas.values())
    assert sum(len(v) for v in entradas.values()) == 67
    assert entradas[18] == []
    assert entradas[1] == ENTRADAS_CLARO
    assert entradas[4] == ENTRADAS_ENTEL
    assert entradas[13] == [nm.EntradaSernac(185, "BANCO SANTANDER CHILE")]
    assert len(entradas[27]) == 3


def test_mapeo_empresa_ausente_devuelve_none(mapeo_real):
    assert nm.Mapeo.cargar(mapeo_real).entradas(fila_de("Empresa nueva", numero=51)) is None


def test_mapeo_respaldo_por_marca_normalizada(escribir):
    path = escribir("mapeo.csv", f"{ENCABEZADO_MAPEO}\n;Banco de Chile;27498;BANCO DE CHILE - BANCO EDWARDS - CITI;\n", bom=True)
    m = nm.Mapeo.cargar(path)
    assert m.entradas(fila_de("BANCO DE CHILE", numero=None)) == [ENTRADA_BANCO_CHILE]
    assert m.entradas(fila_de("Banco de Chile", numero=11)) == [ENTRADA_BANCO_CHILE]


def test_mapeo_detecta_delimitador_coma_con_campos_entre_comillas(escribir):
    texto = (
        "#,Marca,ID SERNAC,Entrada SERNAC,Nota\n"
        '22,Tarjeta Cencosud Scotiabank,216,"CAT/TARJETAS JUMBO, PARIS, EASY, T.BUS, CENCOSUD VISA, MASTERCARD, BANCO PARIS TARJ. Y CRED.CONSUMO",\n'
        f"1,Claro Chile,178,{TEXTO_POR_ID[178]},\n"
    )
    m = nm.Mapeo.cargar(escribir("mapeo.csv", texto))
    assert m.entradas(fila_de("x", numero=22)) == [
        nm.EntradaSernac(216, "CAT/TARJETAS JUMBO, PARIS, EASY, T.BUS, CENCOSUD VISA, MASTERCARD, BANCO PARIS TARJ. Y CRED.CONSUMO")
    ]
    assert m.entradas(fila_de("x", numero=1)) == [ENTRADAS_CLARO[0]]


@pytest.mark.parametrize("bom", [True, False])
def test_mapeo_lee_con_y_sin_bom(escribir, bom):
    texto = f"{ENCABEZADO_MAPEO}\n4;Entel;183;{TEXTO_POR_ID[183]};\n4;Entel;184;{TEXTO_POR_ID[184]};\n"
    m = nm.Mapeo.cargar(escribir("mapeo.csv", texto, bom=bom))
    assert m.entradas(fila_de("Entel", numero=4)) == ENTRADAS_ENTEL[:2]


def test_mapeo_entrada_sin_id_queda_con_id_none(escribir):
    path = escribir("mapeo.csv", f"{ENCABEZADO_MAPEO}\n5;WOM Chile;;WOM SPA (EX-NEXTEL) - VENTAS INTERNET;\n")
    assert nm.Mapeo.cargar(path).entradas(fila_de("WOM Chile", numero=5)) == [nm.EntradaSernac(None, "WOM SPA (EX-NEXTEL) - VENTAS INTERNET")]


@pytest.mark.parametrize(
    "texto",
    [
        "#;Marca;Entrada SERNAC\n1;Claro Chile;CLARO MOVIL\n",
        f"{ENCABEZADO_MAPEO}\n1;Claro Chile;abc;CLARO MOVIL;\n",
    ],
)
def test_mapeo_mal_formado_lanza_error_mapeo(escribir, texto):
    with pytest.raises(nm.ErrorMapeo):
        nm.Mapeo.cargar(escribir("mapeo.csv", texto))


def test_mapeo_inexistente_queda_vacio_con_advertencia(tmp_path, caplog, capsys):
    caplog.set_level(logging.WARNING)
    m = nm.Mapeo.cargar(tmp_path / "no-existe.csv")
    salida = capsys.readouterr()
    assert m.entradas(fila_de("Claro Chile", numero=1)) is None
    assert "no-existe.csv" in caplog.text + salida.out + salida.err


def test_mapeo_registrar_reemplaza_solo_esa_empresa_y_conserva_las_demas(mapeo_real):
    antes = _filas_csv(mapeo_real)
    m = nm.Mapeo.cargar(mapeo_real)
    fila = fila_de("Banco Security", numero=18)
    nuevas = [nm.EntradaSernac(None, "BANCO SECURITY (PRUEBA)")]
    m.registrar(fila, nuevas, nota="confirmado a mano")
    despues = _filas_csv(mapeo_real)
    otras = lambda filas: [tuple(r.values()) for r in filas if r["#"] != "18"]  # noqa: E731
    assert otras(despues) == otras(antes)
    assert nm.Mapeo.cargar(mapeo_real).entradas(fila) == nuevas
    assert any("confirmado a mano" in r["Nota"] for r in despues if r["#"] == "18")
    datos = mapeo_real.read_bytes()
    assert datos.startswith(BOM)
    assert datos.decode("utf-8-sig").splitlines()[0] == ENCABEZADO_MAPEO


def test_mapeo_registrar_lista_vacia_marca_sin_entrada(mapeo_real):
    m = nm.Mapeo.cargar(mapeo_real)
    fila = fila_de("Claro Chile", numero=1)
    m.registrar(fila, [])
    assert nm.Mapeo.cargar(mapeo_real).entradas(fila) == []
    assert [r for r in _filas_csv(mapeo_real) if r["#"] == "1"][0]["ID SERNAC"] == ""


def test_mapeo_registrar_crea_el_archivo_si_no_existe(tmp_path, monkeypatch):
    path = tmp_path / "nuevo.csv"
    real = os.replace
    destinos = []
    monkeypatch.setattr(os, "replace", lambda s, d, *a, **k: (destinos.append(Path(d)), real(s, d, *a, **k))[1])
    m = nm.Mapeo.cargar(path)
    m.registrar(fila_de("Entel", numero=4), ENTRADAS_ENTEL)
    assert path.read_bytes().decode("utf-8-sig").splitlines()[0] == ENCABEZADO_MAPEO
    assert nm.Mapeo.cargar(path).entradas(fila_de("Entel", numero=4)) == ENTRADAS_ENTEL
    assert any(d.resolve() == path.resolve() for d in destinos)


# =============================================================================
# resolver_entradas (mapeo interactivo, con un buscador falso)
# =============================================================================
def _lineas_con(texto: str, contenido: str) -> list[str]:
    return [ln for ln in texto.splitlines() if contenido in ln]


def test_resolver_entradas_enter_acepta_la_sugerencia(tmp_path, capsys):
    mapeo_path = tmp_path / "mapeo.csv"
    m = nm.Mapeo.cargar(mapeo_path)
    bot = BuscadorFalso()
    preguntar = PreguntarGuionado([""])
    r = nm.resolver_entradas(fila_real(11), bot, m, preguntar)
    salida = capsys.readouterr()
    listado = salida.out + salida.err + preguntar.texto
    assert r == [nm.EntradaSernac(None, "BANCO DE CHILE - BANCO EDWARDS - CITI")]
    assert nm.Mapeo.cargar(mapeo_path).entradas(fila_real(11)) == r
    assert any(_norm(b) == "banco de chile" for b in bot.busquedas)
    assert all("*" in ln for ln in _lineas_con(listado, "BANCO DE CHILE - BANCO EDWARDS - CITI"))
    assert _lineas_con(listado, "AFUBACH") and not any("*" in ln for ln in _lineas_con(listado, "AFUBACH"))


def test_resolver_entradas_enter_acepta_varias_sugerencias(tmp_path):
    m = nm.Mapeo.cargar(tmp_path / "mapeo.csv")
    r = nm.resolver_entradas(fila_real(3), BuscadorFalso(), m, PreguntarGuionado([""]))
    assert {e.texto for e in r} == set(_textos("movistar"))
    assert all(e.id is None for e in r)


def test_resolver_entradas_numeros_eligen_de_la_lista_ordenada(tmp_path):
    m = nm.Mapeo.cargar(tmp_path / "mapeo.csv")
    r = nm.resolver_entradas(fila_real(4), BuscadorFalso(), m, PreguntarGuionado(["1,2,3"]))
    assert {e.texto for e in r} == {TEXTO_POR_ID[183], TEXTO_POR_ID[184], TEXTO_POR_ID[432]}


def test_resolver_entradas_b_busca_otro_texto_y_n_registra_sin_entrada(tmp_path, capsys):
    mapeo_path = tmp_path / "mapeo.csv"
    m = nm.Mapeo.cargar(mapeo_path)
    bot = BuscadorFalso(resultados={"security": ["SEGURIDAD SECURITY LTDA", "INVERSIONES SECURITY CENTER"]})
    r = nm.resolver_entradas(fila_real(18), bot, m, PreguntarGuionado(["b security", "n"]))
    salida = capsys.readouterr()
    assert r == []
    assert "security" in [_norm(b) for b in bot.busquedas]
    assert "SEGURIDAD SECURITY LTDA" in salida.out + salida.err
    assert nm.Mapeo.cargar(mapeo_path).entradas(fila_real(18)) == []


def test_resolver_entradas_b_y_luego_numero(tmp_path):
    m = nm.Mapeo.cargar(tmp_path / "mapeo.csv")
    bot = BuscadorFalso(resultados={"vtr telefonia": ["VTR TELEFONIA MOVIL Y FIJO (HOGAR)"]})
    fila = fila_de("VTR", numero=2)
    r = nm.resolver_entradas(fila, bot, m, PreguntarGuionado(["b vtr telefonia", "1"]))
    assert r == [nm.EntradaSernac(None, "VTR TELEFONIA MOVIL Y FIJO (HOGAR)")]


def test_resolver_entradas_s_omite_sin_registrar(tmp_path):
    mapeo_path = tmp_path / "mapeo.csv"
    m = nm.Mapeo.cargar(mapeo_path)
    r = nm.resolver_entradas(fila_real(1), BuscadorFalso(), m, PreguntarGuionado(["s"]))
    assert r is None
    assert m.entradas(fila_real(1)) is None
    assert not mapeo_path.exists()


def test_resolver_entradas_muestra_como_maximo_25_opciones(tmp_path, capsys):
    m = nm.Mapeo.cargar(tmp_path / "mapeo.csv")
    bot = BuscadorFalso(resultados={"opcion": [f"OPCION {i:02d}" for i in range(1, 31)]})
    preguntar = PreguntarGuionado(["s"])
    nm.resolver_entradas(fila_de("Opcion", numero=7), bot, m, preguntar)
    salida = capsys.readouterr()
    listado = salida.out + salida.err + preguntar.texto
    assert "OPCION 25" in listado
    assert "OPCION 26" not in listado


# =============================================================================
# Correo (FakeSMTP)
# =============================================================================
def crear_enviador(persona, fabrica, *, sleep=None, usuario=GMAIL_USUARIO, tiempos=None):
    return nm.EnviadorCorreo(
        usuario,
        nm.Secreto(SECRETO_GMAIL_SIN_ESPACIOS),
        persona,
        smtp_factory=fabrica,
        sleep=sleep if sleep is not None else SleepGrabador(),
        tiempos=tiempos or tiempos_prueba(),
    )


def test_enviador_conectar_usa_gmail_ssl_y_hace_login(persona):
    fab = FabricaSMTP()
    crear_enviador(persona, fab).conectar()
    host, puerto, kwargs = fab.conexiones[0]
    assert (host, puerto) == ("smtp.gmail.com", 465)
    assert isinstance(kwargs.get("context"), ssl.SSLContext)
    assert kwargs.get("timeout") == 30
    assert fab.logins == [(GMAIL_USUARIO, SECRETO_GMAIL_SIN_ESPACIOS)]


def test_enviador_construir_mensaje_con_cabeceras_completas(persona):
    msg = crear_enviador(persona, FabricaSMTP()).construir_mensaje("protecciondedatos@entel.cl", "Oposición – Entel", "Hola Entel,\nTexto ñandú.")
    assert isinstance(msg, EmailMessage)
    assert parseaddr(str(msg["From"])) == ("Ana Pérez Soto", GMAIL_USUARIO)
    assert str(msg["To"]) == "protecciondedatos@entel.cl"
    assert str(msg["Reply-To"]) == CORREO_PERSONA
    assert str(msg["Subject"]) == "Oposición – Entel"
    assert msg["Date"] and msg["Message-ID"]
    assert msg.get_content_type() == "text/plain"
    assert msg.get_content_charset() == "utf-8"
    assert "Texto ñandú." in msg.get_content()


def test_enviador_sin_reply_to_si_el_usuario_es_el_correo_personal(persona):
    msg = crear_enviador(persona, FabricaSMTP(), usuario=CORREO_PERSONA).construir_mensaje("c@entel.cl", "A", "B")
    assert msg["Reply-To"] is None


@pytest.mark.parametrize(
    "destinatario, asunto",
    [("c@entel.cl\r\nBcc: espia@malo.cl", "Asunto"), ("c@entel.cl", "Hola\nBcc: espia@malo.cl"), ("c@entel.cl", "Hola\rX")],
)
def test_enviador_rechaza_saltos_de_linea_en_cabeceras(persona, destinatario, asunto):
    with pytest.raises(ValueError):
        crear_enviador(persona, FabricaSMTP()).construir_mensaje(destinatario, asunto, "cuerpo")


def test_enviador_reconecta_y_reintenta_si_se_corta_la_conexion(persona, sleep_falso):
    fab = FabricaSMTP({"send_message": [smtplib.SMTPServerDisconnected("se cortó")]})
    crear_enviador(persona, fab, sleep=sleep_falso).enviar("c@entel.cl", "A", "B")
    assert len(fab.enviados) == 1
    assert len(fab.conexiones) == 2
    assert len(fab.logins) == 2
    assert sleep_falso.llamadas


def test_enviador_reintenta_respuestas_4xx(persona):
    fab = FabricaSMTP({"send_message": [smtplib.SMTPDataError(451, b"4.3.0 intente mas tarde")]})
    crear_enviador(persona, fab).enviar("c@entel.cl", "A", "B")
    assert fab.intentos_envio == 2
    assert len(fab.enviados) == 1


@pytest.mark.parametrize(
    "exc",
    [smtplib.SMTPRecipientsRefused({"c@entel.cl": (550, b"5.1.1 no existe")}), smtplib.SMTPDataError(550, b"5.7.1 rechazado")],
)
def test_enviador_no_reintenta_errores_permanentes(persona, exc):
    fab = FabricaSMTP({"send_message": [exc]})
    with pytest.raises(type(exc)):
        crear_enviador(persona, fab).enviar("c@entel.cl", "A", "B")
    assert fab.intentos_envio == 1
    assert fab.enviados == []


def test_enviador_agota_reintentos_transitorios(persona):
    fab = FabricaSMTP({"send_message": [smtplib.SMTPServerDisconnected("cortado")] * 10})
    with pytest.raises(smtplib.SMTPServerDisconnected):
        crear_enviador(persona, fab).enviar("c@entel.cl", "A", "B")
    assert fab.intentos_envio == tiempos_prueba()["reintentos"]


PLANTILLA_SIMPLE = nm.Plantilla("Hola {empresa}", "Estimados de {empresa}, soy {nombre}. Sector: {sector}.")


def _tracker_correo(fabrica_tracker, filas):
    return fabrica_tracker([{"Vía de envío": "Correo", "Sector": "Retail", **f} for f in filas])


def test_enviador_procesar_envia_uno_por_empresa_y_solo_escribe_enviado_fecha(persona, fabrica_tracker, sleep_falso):
    path = _tracker_correo(
        fabrica_tracker,
        [
            {"Marca": "Paris (Cencosud retail)", "Correo datos personales": "datospersonales@cencosud.cl"},
            {"Marca": "Falabella retail (falabella.com)", "Correo datos personales": "tusdatospersonales@falabella.cl"},
            {"Marca": "Lider / Walmart Chile", "Correo datos personales": "servicioalcliente@lider.cl"},
        ],
    )
    t = nm.Tracker(path)
    fab = FabricaSMTP()
    resultados = crear_enviador(persona, fab, sleep=sleep_falso).procesar(t.filas(), PLANTILLA_SIMPLE, t, hoy=HOY, preguntar=PreguntarGuionado())
    assert [(r.numero_fila, r.fase, r.estado) for r in resultados] == [(2, "correo", "ok"), (3, "correo", "ok"), (4, "correo", "ok")]
    assert [str(m["To"]) for m in fab.enviados] == ["datospersonales@cencosud.cl", "tusdatospersonales@falabella.cl", "servicioalcliente@lider.cl"]
    assert str(fab.enviados[0]["Subject"]) == "Hola Paris (Cencosud retail)"
    assert "Estimados de Paris (Cencosud retail), soy Ana Pérez Soto. Sector: Retail." in fab.enviados[0].get_content()
    assert len(fab.conexiones) == 1
    for fila in (2, 3, 4):
        assert es_fecha(valor_celda(path, fila, "Enviado (fecha)"), HOY)
        assert vacio(valor_celda(path, fila, "No Molestar ingresado"))
        assert str(valor_celda(path, fila, "Estado")).startswith("=")
    assert len(sleep_falso.en_rango(70, 80)) == 2


def test_enviador_procesar_sin_correo_valido_marca_fallida_y_no_toca_el_tracker(persona, fabrica_tracker):
    path = _tracker_correo(
        fabrica_tracker,
        [
            {"Marca": "Paris (Cencosud retail)", "Correo datos personales": "datospersonales@cencosud.cl"},
            {"Marca": "Sin Correo"},
            {"Marca": "Mal Correo", "Correo datos personales": "no-es-correo"},
        ],
    )
    t = nm.Tracker(path)
    fab = FabricaSMTP()
    resultados = crear_enviador(persona, fab).procesar(t.filas(), PLANTILLA_SIMPLE, t, hoy=HOY, preguntar=PreguntarGuionado())
    assert [r.estado for r in resultados] == ["ok", "fallida", "fallida"]
    assert all("Falta correo válido de la empresa" in r.detalle for r in resultados[1:])
    assert len(fab.enviados) == 1
    for fila in (3, 4):
        assert vacio(valor_celda(path, fila, "Enviado (fecha)"))
    texto = todas_las_celdas(path)
    assert "Falta correo" not in texto and "Error" not in texto


def test_enviador_procesar_error_permanente_sigue_con_las_demas(persona, fabrica_tracker):
    path = _tracker_correo(
        fabrica_tracker,
        [
            {"Marca": "Paris (Cencosud retail)", "Correo datos personales": "datospersonales@cencosud.cl"},
            {"Marca": "Lider / Walmart Chile", "Correo datos personales": "servicioalcliente@lider.cl"},
        ],
    )
    t = nm.Tracker(path)
    fab = FabricaSMTP({"send_message": [smtplib.SMTPRecipientsRefused({"datospersonales@cencosud.cl": (550, b"no existe")}), None]})
    resultados = crear_enviador(persona, fab).procesar(t.filas(), PLANTILLA_SIMPLE, t, hoy=HOY, preguntar=PreguntarGuionado())
    assert [r.estado for r in resultados] == ["fallida", "ok"]
    assert [str(m["To"]) for m in fab.enviados] == ["servicioalcliente@lider.cl"]
    assert vacio(valor_celda(path, 2, "Enviado (fecha)"))
    assert es_fecha(valor_celda(path, 3, "Enviado (fecha)"), HOY)


def test_enviador_procesar_error_de_autenticacion_aborta_la_fase(persona, fabrica_tracker):
    path = _tracker_correo(fabrica_tracker, [{"Marca": "Paris (Cencosud retail)", "Correo datos personales": "datospersonales@cencosud.cl"}])
    t = nm.Tracker(path)
    fab = FabricaSMTP({"login": [smtplib.SMTPAuthenticationError(535, b"5.7.8 Username and Password not accepted")]})
    with pytest.raises(smtplib.SMTPAuthenticationError):
        crear_enviador(persona, fab).procesar(t.filas(), PLANTILLA_SIMPLE, t, hoy=HOY, preguntar=PreguntarGuionado())
    assert fab.enviados == []
    assert len(fab.logins) == 1


# =============================================================================
# BotSernac con FakePage (página + iframe #frameNoMolestar con su propio DOM)
# =============================================================================
def test_abrir_portal_navega_a_portal_url(nuevo_portal, crear_bot):
    p = nuevo_portal()
    crear_bot(p.page).abrir_portal()
    assert p.top.gotos == [nm.PORTAL_URL]


@pytest.mark.parametrize("status", [403, 500, 503])
def test_abrir_portal_con_403_o_5xx_lanza_error_portal_con_pista_de_mantencion(nuevo_portal, crear_bot, status):
    p = nuevo_portal(status_portal=status)
    with pytest.raises(nm.ErrorPortal) as e:
        crear_bot(p.page).abrir_portal()
    assert "mantenci" in str(e.value).lower()


def _en_login(nuevo_portal, **opciones) -> PortalFalso:
    p = nuevo_portal(**opciones)
    p.page.goto(nm.PORTAL_URL)
    return p


def test_iniciar_sesion_tipea_con_press_sequentially_y_nunca_fill(nuevo_portal, crear_bot):
    p = _en_login(nuevo_portal)
    crear_bot(p.page).iniciar_sesion(nm.Secreto(SECRETO_CU))
    assert p.logueado
    assert p.envios_login == [("123456785", SECRETO_CU)]
    run, clave, ingresar = p.top_cu["run"], p.top_cu["clave"], p.top_cu["ingresar"]
    assert run.llenados == [] and clave.llenados == []
    assert "".join(run.escrito) == "123456785"
    assert "".join(clave.escrito) == SECRETO_CU
    assert ingresar.clicks == 1
    assert p.top.capturas == []


def test_iniciar_sesion_respeta_formato_run(nuevo_portal, crear_bot, monkeypatch):
    monkeypatch.setattr(nm, "FORMATO_RUN", "sin_puntos")
    p = _en_login(nuevo_portal)
    crear_bot(p.page).iniciar_sesion(nm.Secreto(SECRETO_CU))
    assert p.envios_login == [("12345678-5", SECRETO_CU)]


def test_iniciar_sesion_con_sesion_ya_iniciada_no_hace_nada(nuevo_portal, crear_bot):
    p = nuevo_portal()
    p.entrar()
    crear_bot(p.page).iniciar_sesion(nm.Secreto(SECRETO_CU))
    assert p.envios_login == []
    assert p.top_cu is None


def _con_logout_oculto_primero(p: PortalFalso) -> None:
    """Como el portal real: dos button[onclick='logoutCU()'], el primero oculto (menú móvil)."""
    home_original = p.home

    def home() -> None:
        home_original()
        p.top.elementos[S("sesion_iniciada")].insert(0, Elemento(texto="", visible=False))

    p.home = home


def test_iniciar_sesion_detecta_la_sesion_aunque_el_primer_logout_este_oculto(nuevo_portal, crear_bot):
    p = nuevo_portal()
    _con_logout_oculto_primero(p)
    p.entrar()
    crear_bot(p.page).iniciar_sesion(nm.Secreto(SECRETO_CU))
    assert p.envios_login == []


def test_iniciar_sesion_detecta_la_sesion_con_el_logout_oculto_en_un_menu(nuevo_portal, crear_bot):
    p = nuevo_portal()
    p.entrar()
    for e in p.top.elementos[S("sesion_iniciada")]:
        e.visible = False  # el botón existe solo con sesión, aunque esté dentro de un menú cerrado
    crear_bot(p.page).iniciar_sesion(nm.Secreto(SECRETO_CU))
    assert p.envios_login == []


def test_iniciar_sesion_termina_al_volver_con_el_primer_logout_oculto(nuevo_portal, crear_bot):
    p = _en_login(nuevo_portal)
    _con_logout_oculto_primero(p)
    crear_bot(p.page).iniciar_sesion(nm.Secreto(SECRETO_CU))
    assert p.logueado
    assert p.envios_login == [("123456785", SECRETO_CU)]


def test_iniciar_sesion_prueba_las_alternativas_de_selector_en_orden(nuevo_portal, crear_bot):
    selectores = dict(nm.SELECTORES, login_boton_claveunica=["#no-existe-boton", "#btnClaveUnica"])
    p = _en_login(nuevo_portal)
    crear_bot(p.page, selectores=selectores).iniciar_sesion(nm.Secreto(SECRETO_CU))
    assert p.logueado


def test_iniciar_sesion_con_error_lanza_error_login_sin_reintentar(nuevo_portal, crear_bot):
    p = _en_login(nuevo_portal, login="error")
    with pytest.raises(nm.ErrorLogin) as e:
        crear_bot(p.page).iniciar_sesion(nm.Secreto(SECRETO_CU))
    assert len(p.envios_login) == 1
    assert "".join(p.top_cu["clave"].escrito) == SECRETO_CU
    assert SECRETO_CU not in str(e.value)


def test_iniciar_sesion_sin_respuesta_agota_el_tiempo_y_no_reintenta(nuevo_portal, crear_bot):
    p = _en_login(nuevo_portal, login="nada")
    with pytest.raises(nm.ErrorLogin):
        crear_bot(p.page, tiempos=tiempos_prueba(timeout_login_s=0.3)).iniciar_sesion(nm.Secreto(SECRETO_CU))
    assert len(p.envios_login) == 1
    assert p.top_cu["ingresar"].clicks == 1


def test_iniciar_sesion_con_2fa_pausa_y_continua(nuevo_portal, crear_bot):
    p = _en_login(nuevo_portal, login="2fa")
    preguntar = PreguntarGuionado([lambda prompt: (p.resolver_2fa(), "")[1]])
    crear_bot(p.page, preguntar=preguntar).iniciar_sesion(nm.Secreto(SECRETO_CU))
    assert preguntar.respuestas == []
    assert p.logueado
    assert len(p.envios_login) == 1


def test_iniciar_sesion_con_captcha_visible_pausa_antes_de_tipear_y_nunca_lo_toca(nuevo_portal, crear_bot, capsys):
    p = _en_login(nuevo_portal, captcha_en_login=True)
    visto = {}

    def persona_resuelve(prompt):
        visto["captchas"] = [e for s in nm.SELECTORES_CAPTCHA for e in p.top.todos(s)]
        visto["clave_tipeada"] = "".join(p.top_cu["clave"].escrito)
        p.resolver_captcha()
        return ""

    preguntar = PreguntarGuionado([persona_resuelve])
    crear_bot(p.page, preguntar=preguntar).iniciar_sesion(nm.Secreto(SECRETO_CU))
    salida = capsys.readouterr()
    assert preguntar.respuestas == []
    assert "captcha" in (preguntar.texto + salida.out + salida.err).lower()
    assert visto["clave_tipeada"] == ""
    assert visto["captchas"] and all(e.clicks == 0 for e in visto["captchas"])
    assert p.logueado
    assert p.envios_login == [("123456785", SECRETO_CU)]


def test_iniciar_sesion_con_captcha_tras_ingresar_pausa_y_no_reenvia_la_clave(nuevo_portal, crear_bot):
    p = _en_login(nuevo_portal, login="captcha")
    preguntar = PreguntarGuionado([lambda prompt: (p.resolver_captcha(entrar=True), "")[1]])
    crear_bot(p.page, preguntar=preguntar).iniciar_sesion(nm.Secreto(SECRETO_CU))
    assert preguntar.respuestas == []
    assert p.logueado
    assert p.top_cu["ingresar"].clicks == 1
    assert len(p.envios_login) == 1


def test_iniciar_sesion_ignora_el_badge_invisible_de_256x60(nuevo_portal, crear_bot):
    p = _en_login(nuevo_portal)
    assert p.top.todos(".grecaptcha-badge") and p.top.el("iframe[src*='size=invisible']").caja["width"] == 256
    crear_bot(p.page, preguntar=PreguntarGuionado()).iniciar_sesion(nm.Secreto(SECRETO_CU))
    assert p.logueado


def test_hay_captcha_false_en_pagina_limpia(nuevo_portal, crear_bot):
    p = nuevo_portal()
    p.entrar()
    assert crear_bot(p.page).hay_captcha() is False


def test_hay_captcha_false_con_solo_el_badge_invisible(nuevo_portal, crear_bot):
    p = _en_login(nuevo_portal)
    assert crear_bot(p.page).hay_captcha() is False


def test_hay_captcha_true_con_recaptcha_v2_visible(nuevo_portal, crear_bot):
    p = _en_login(nuevo_portal)
    p.top.frames.append(URL_ANCHOR_V2)
    marcar_captcha(p.top, ("anchor",))
    assert crear_bot(p.page).hay_captcha() is True


def test_hay_captcha_false_con_bframe_oculto(nuevo_portal, crear_bot):
    p = _en_login(nuevo_portal)
    p.top.frames.append(URL_BFRAME)
    marcar_captcha(p.top, ("bframe",), visible=False, caja=BBOX_NULO)
    assert crear_bot(p.page).hay_captcha() is False


def test_hay_captcha_true_con_desafio_bframe_visible(nuevo_portal, crear_bot):
    p = _en_login(nuevo_portal)
    p.top.frames.append(URL_BFRAME)
    marcar_captcha(p.top, ("bframe",), visible=True, caja={"x": 100.0, "y": 50.0, "width": 400.0, "height": 580.0})
    assert crear_bot(p.page).hay_captcha() is True


def test_hay_captcha_true_con_hcaptcha(nuevo_portal, crear_bot):
    p = _en_login(nuevo_portal)
    p.top.frames.append(URL_HCAPTCHA)
    marcar_captcha(p.top, ("hcaptcha", "h-captcha"))
    assert crear_bot(p.page).hay_captcha() is True


def test_hay_captcha_true_con_turnstile(nuevo_portal, crear_bot):
    p = _en_login(nuevo_portal)
    p.top.frames.append(URL_TURNSTILE)
    marcar_captcha(p.top, ("cloudflare",))
    assert crear_bot(p.page).hay_captcha() is True


def test_hay_captcha_true_con_texto_no_soy_un_robot(nuevo_portal, crear_bot):
    p = _en_login(nuevo_portal)
    p.top.texto = "Antes de continuar: No soy un robot"
    marcar_captcha(p.top, ("robot",))
    assert crear_bot(p.page).hay_captcha() is True


def test_hay_2fa_true_con_campo_de_codigo_visible(nuevo_portal, crear_bot):
    p = _en_login(nuevo_portal)
    p.top.agregar(nm.SELECTORES_2FA[0], visible=True)
    assert crear_bot(p.page).hay_2fa() is True


@pytest.mark.parametrize("texto", ["Ingresa el código de verificación enviado a tu celular", "Activa tu segundo factor", "Código de seguridad"])
def test_hay_2fa_true_por_texto_de_la_pagina(nuevo_portal, crear_bot, texto):
    p = _en_login(nuevo_portal)
    p.top.texto = texto
    assert crear_bot(p.page).hay_2fa() is True


def test_hay_2fa_false_en_login_normal_o_campo_oculto(nuevo_portal, crear_bot):
    p = _en_login(nuevo_portal)
    p.top.texto = "Ingresa tu RUN y tu ClaveÚnica"
    p.top.agregar(nm.SELECTORES_2FA[0], visible=False)
    assert crear_bot(p.page).hay_2fa() is False


def test_mantener_sesion_hace_clic_si_el_modal_esta_visible(nuevo_portal, crear_bot):
    p = nuevo_portal(modal_sesion=True)
    p.entrar()
    crear_bot(p.page).mantener_sesion()
    assert p.clicks_mantener == 1
    assert not p.top.el(S("boton_mantener_sesion")).visible


def test_mantener_sesion_sin_modal_no_hace_nada(nuevo_portal, crear_bot):
    p = nuevo_portal()
    p.entrar()
    crear_bot(p.page).mantener_sesion()
    assert p.clicks_mantener == 0


def test_buscar_empresas_abre_el_formulario_si_hace_falta_y_devuelve_textos(nuevo_portal, crear_bot):
    p = nuevo_portal()
    p.entrar()
    sleep = SleepGrabador()
    textos = crear_bot(p.page, sleep=sleep).buscar_empresas("CLARO")
    assert textos == [t for _, t in buscar_en_directorio("CLARO")]
    assert p.busquedas[-1] == "CLARO"
    assert p.clicks_nueva == 1
    assert len(sleep.en_rango(0.003, 0.004)) == 1


def test_buscar_empresas_con_formulario_abierto_no_vuelve_a_navegar(nuevo_portal, crear_bot):
    p = nuevo_portal()
    p.entrar()
    p.page.goto(nm.URL_NO_MOLESTAR)
    p.formulario()
    gotos = len(p.top.gotos)
    assert crear_bot(p.page).buscar_empresas("BANCO DE CHILE") == _textos("banco de chile")
    assert len(p.top.gotos) == gotos


def test_buscar_empresas_sin_resultados_devuelve_lista_vacia(nuevo_portal, crear_bot):
    p = nuevo_portal()
    p.entrar()
    assert crear_bot(p.page).buscar_empresas("ZZZ EMPRESA INEXISTENTE") == []


def test_buscar_empresas_colapsa_espacios_de_las_opciones(nuevo_portal, crear_bot):
    p = nuevo_portal(directorio=[(1, "BANCO  DE\nCHILE   EDWARDS ")])
    p.entrar()
    assert crear_bot(p.page).buscar_empresas("BANCO") == ["BANCO DE CHILE EDWARDS"]


def _listo(nuevo_portal, **opciones) -> PortalFalso:
    p = nuevo_portal(**opciones)
    p.entrar()
    return p


def test_enviar_solicitud_claro_dos_entradas_ok_con_numero(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal)
    r = crear_bot(p.page).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert (r.numero_fila, r.empresa, r.fase, r.estado, r.numero) == (2, "Claro Chile", "sernac", "ok", "1234567")
    assert p.envios == 1
    sol = p.solicitudes[0]
    assert sol["empresas"] == [(178, TEXTO_POR_ID[178]), (146420, TEXTO_POR_ID[146420])]
    assert sol["telefonos"] == [TEL_PERSONA_E164]
    assert sol["correos"] == [CORREO_PERSONA]
    assert len(sol["canales_tel"]) == 6 and all(sol["canales_tel"])
    assert len(sol["canales_mail"]) == 2 and all(sol["canales_mail"])
    assert sol["acepto"] is True
    assert "CLARO MOVIL" in p.busquedas and "CLARO FIJO" in p.busquedas


def test_enviar_solicitud_no_duplica_el_telefono_precargado_y_agrega_el_correo_por_partes(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal)
    crear_bot(p.page).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert p.contactos_agregados == [("correo", CORREO_PERSONA, None)]


def test_enviar_solicitud_agrega_el_telefono_movil_sin_quitar_el_precargado(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal, telefonos_precargados=("+56987654321",))
    r = crear_bot(p.page).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "ok"
    assert ("telefono", TEL_PERSONA_E164, "Móvil") in p.contactos_agregados
    assert p.solicitudes[0]["telefonos"] == ["+56987654321", TEL_PERSONA_E164]


def test_enviar_solicitud_telefono_fijo_elige_tipo_fijo(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal, telefonos_precargados=())
    fijo = nm.Persona(nombre="Ana Pérez Soto", rut="12345678-5", telefono="223456789", correo=CORREO_PERSONA)
    r = crear_bot(p.page, persona_bot=fijo).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "ok"
    assert ("telefono", "+56223456789", "Fijo") in p.contactos_agregados


def test_enviar_solicitud_no_duplica_el_correo_precargado_aunque_cambien_mayusculas(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal, correos_precargados=("Ana.Perez@Example.cl",))
    r = crear_bot(p.page).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "ok"
    assert p.contactos_agregados == []
    assert p.solicitudes[0]["correos"] == ["Ana.Perez@Example.cl"]


def test_enviar_solicitud_error_del_portal_al_agregar_contacto_queda_fallida(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal, telefonos_precargados=(), rechazar_telefono=True)
    r = crear_bot(p.page).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "fallida"
    assert "teléfono válido" in r.detalle
    assert p.envios == 0


def test_enviar_solicitud_marca_el_consentimiento_si_viene_desmarcado(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal, acepto_marcado=False)
    r = crear_bot(p.page).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "ok"
    assert p.solicitudes[0]["acepto"] is True


def test_enviar_solicitud_hace_clic_en_mantener_sesion_si_aparece(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal, modal_sesion=True)
    r = crear_bot(p.page).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "ok"
    assert p.clicks_mantener >= 1


def test_enviar_solicitud_marca_todos_los_canales_de_telefono_y_correo(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal, canales_tel_desmarcados=(1, 4), canales_mail_desmarcados=(0,))
    r = crear_bot(p.page).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "ok"
    assert all(p.solicitudes[0]["canales_tel"]) and all(p.solicitudes[0]["canales_mail"])


def test_enviar_solicitud_sigue_continuar_confirmar_y_finalizar_en_orden(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal)
    crear_bot(p.page).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert [e for e in p.eventos if isinstance(e, str)] == ["continuar", "confirmar", "finalizar"]


def test_enviar_solicitud_con_confirmacion_s_envia(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal)
    preguntar = PreguntarGuionado(["s"])
    r = crear_bot(p.page, preguntar=preguntar, confirmar=True).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "ok"
    prompt = preguntar.prompts[0]
    assert "Claro Chile" in prompt and "2" in prompt and "[s/N]" in prompt
    assert p.envios == 1


def test_enviar_solicitud_con_confirmacion_n_omite_sin_finalizar(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal)
    r = crear_bot(p.page, preguntar=PreguntarGuionado(["n"]), confirmar=True).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "omitida"
    assert p.envios == 0
    assert "finalizar" not in p.eventos


def _opciones_de(termino: str) -> list[str]:
    return [t for _, t in buscar_en_directorio(termino)]


def test_enviar_solicitud_sin_coincidencia_exacta_elige_por_numero(nuevo_portal, crear_bot, capsys):
    p = _listo(nuevo_portal)
    numero = _opciones_de("FALABELLA").index("TIENDAS FALABELLA") + 1
    preguntar = PreguntarGuionado([str(numero)])
    fila = fila_de("Falabella retail (falabella.com)", numero=28)
    r = crear_bot(p.page, preguntar=preguntar).enviar_solicitud(fila, [nm.EntradaSernac(None, "FALABELLA")])
    salida = capsys.readouterr()
    assert r.estado == "ok"
    assert p.solicitudes[0]["empresas"] == [(182, "TIENDAS FALABELLA")]
    listado = preguntar.texto + salida.out + salida.err
    assert "TIENDAS FALABELLA" in listado and "FALABELLA.COM - LINIO" in listado


def test_enviar_solicitud_sin_coincidencia_m_usa_lo_elegido_en_el_navegador(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal)
    preguntar = PreguntarGuionado([lambda prompt: (p.elegir_manual(182), "m")[1]])
    fila = fila_de("Falabella retail (falabella.com)", numero=28)
    r = crear_bot(p.page, preguntar=preguntar).enviar_solicitud(fila, [nm.EntradaSernac(None, "FALABELLA")])
    assert r.estado == "ok"
    assert p.solicitudes[0]["empresas"] == [(182, "TIENDAS FALABELLA")]


def test_enviar_solicitud_s_omite_una_entrada_y_envia_las_demas(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal)
    entradas = [nm.EntradaSernac(None, "CLARO TELEFONIA INEXISTENTE"), ENTRADAS_CLARO[0]]
    r = crear_bot(p.page, preguntar=PreguntarGuionado(["s"])).enviar_solicitud(fila_real(1), entradas)
    assert r.estado == "ok"
    assert p.solicitudes[0]["empresas"] == [(178, TEXTO_POR_ID[178])]


def test_enviar_solicitud_todas_las_entradas_omitidas_queda_sin_coincidencia(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal)
    r = crear_bot(p.page, preguntar=PreguntarGuionado(["s"])).enviar_solicitud(fila_real(1), [nm.EntradaSernac(None, "EMPRESA FANTASMA SPA")])
    assert r.estado == "sin_coincidencia"
    assert p.clicks_continuar == 0
    assert p.envios == 0


def test_enviar_solicitud_id_distinto_al_mapeo_se_trata_como_sin_coincidencia(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal)
    preguntar = PreguntarGuionado(["s"])
    r = crear_bot(p.page, preguntar=preguntar).enviar_solicitud(fila_real(1), [nm.EntradaSernac(999, "TIENDAS FALABELLA")])
    assert preguntar.respuestas == []
    assert r.estado == "sin_coincidencia"
    assert p.empresas_agregadas == []
    assert p.envios == 0


def test_enviar_solicitud_reintenta_con_el_texto_completo_si_el_termino_no_basta(nuevo_portal, crear_bot):
    directorio = [
        (57701, "AFUBACH (ASOCIACION DE FUNCIONARIOS DEL BANCO DE CHILE)"),
        (900001, "SINDICATO BANCO DE CHILE"),
        (900002, "CLUB DEPORTIVO BANCO DE CHILE"),
        (27498, "BANCO DE CHILE - BANCO EDWARDS - CITI"),
    ]
    p = _listo(nuevo_portal, directorio=directorio, limite_resultados=2)
    r = crear_bot(p.page, preguntar=PreguntarGuionado()).enviar_solicitud(fila_real(11), [ENTRADA_BANCO_CHILE])
    assert r.estado == "ok"
    assert "BANCO DE CHILE" in p.busquedas
    assert "BANCO DE CHILE - BANCO EDWARDS - CITI" in p.busquedas
    assert p.solicitudes[0]["empresas"] == [(27498, "BANCO DE CHILE - BANCO EDWARDS - CITI")]


def test_enviar_solicitud_confirmacion_no_detectada_o_marca_incierta_y_finaliza_una_vez(nuevo_portal, crear_bot, tmp_path):
    p = _listo(nuevo_portal, resultados_envio=(("nada",),))
    capturas = tmp_path / "salida" / "capturas"
    r = crear_bot(p.page, preguntar=PreguntarGuionado(["o"]), capturas_dir=capturas).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "incierta"
    assert p.envios == 1
    assert r.captura and Path(r.captura).exists()
    assert len(list(capturas.glob("*.png"))) == 2


def test_enviar_solicitud_confirmacion_no_detectada_s_con_numero_manual(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal, resultados_envio=(("nada",),))
    r = crear_bot(p.page, preguntar=PreguntarGuionado(["s", "7654321"])).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert (r.estado, r.numero) == ("ok", "7654321")
    assert p.envios == 1


def test_enviar_solicitud_confirmacion_no_detectada_s_sin_numero(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal, resultados_envio=(("nada",),))
    r = crear_bot(p.page, preguntar=PreguntarGuionado(["s", ""])).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert (r.estado, r.numero) == ("ok", None)
    assert p.envios == 1


def test_enviar_solicitud_confirmacion_no_detectada_eof_marca_incierta(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal, resultados_envio=(("nada",),))
    r = crear_bot(p.page, preguntar=PreguntarGuionado([EOFError])).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "incierta"
    assert p.envios == 1


def test_enviar_solicitud_n_permite_un_solo_intento_completo_mas(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal, resultados_envio=(("nada",), ("ok", "1234567")))
    r = crear_bot(p.page, preguntar=PreguntarGuionado(["n"])).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert (r.estado, r.numero) == ("ok", "1234567")
    assert p.envios == 2
    assert p.gotos_no_molestar >= 2


def test_enviar_solicitud_n_dos_veces_queda_fallida_con_maximo_dos_envios(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal, resultados_envio=(("nada",),))
    r = crear_bot(p.page, preguntar=PreguntarGuionado(["n", "n"])).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "fallida"
    assert p.envios == 2


def test_enviar_solicitud_error_al_hacer_clic_en_finalizar_no_reintenta_solo(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal, resultados_envio=(("error_click",),))
    r = crear_bot(p.page, preguntar=PreguntarGuionado(["o"])).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "incierta"
    assert p.envios == 1


def test_enviar_solicitud_exito_por_listado_sin_alerta(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal, resultados_envio=(("listado",),))
    r = crear_bot(p.page).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert (r.estado, r.numero) == ("ok", None)
    assert r.detalle == nm.DETALLE_SIN_NUMERO


def test_enviar_solicitud_busca_el_numero_en_el_texto_del_marco(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal, resultados_envio=(("folio_en_pagina", "98765"),))
    r = crear_bot(p.page).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert (r.estado, r.numero) == ("ok", "98765")


def test_enviar_solicitud_confirmada_sin_numero_explica_que_sernac_no_lo_muestra(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal, resultados_envio=(("sin_numero",),))
    r = crear_bot(p.page).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert (r.estado, r.numero) == ("ok", None)
    assert r.detalle == nm.DETALLE_SIN_NUMERO
    assert r.captura  # la captura de la confirmación sigue guardándose


def test_enviar_solicitud_reintenta_la_preparacion_con_url_no_molestar_fresca(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal)
    p.fallos_nueva_solicitud = [PWTimeoutError("portal lento")]
    sleep = SleepGrabador()
    r = crear_bot(p.page, sleep=sleep).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "ok"
    assert p.clicks_nueva == 2
    assert p.gotos_no_molestar >= 2
    assert p.envios == 1
    assert p.solicitudes[0]["empresas"] == [(178, TEXTO_POR_ID[178]), (146420, TEXTO_POR_ID[146420])]
    assert any(s >= tiempos_prueba()["backoff_base_s"] for s in sleep.llamadas)


def test_enviar_solicitud_preparacion_agotada_queda_fallida_sin_finalizar(nuevo_portal, crear_bot):
    p = _listo(nuevo_portal)
    p.fallos_nueva_solicitud = [PWTimeoutError("portal caído")] * 10
    r = crear_bot(p.page).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "fallida"
    assert p.clicks_nueva == tiempos_prueba()["reintentos"]
    assert p.envios == 0


def test_enviar_solicitud_guarda_captura_de_pagina_y_del_marco(nuevo_portal, crear_bot, tmp_path):
    p = _listo(nuevo_portal)
    capturas = tmp_path / "salida" / "capturas"
    r = crear_bot(p.page, capturas_dir=capturas).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    archivos = sorted(capturas.glob("*.png"))
    assert len(archivos) == 2
    assert all(re.fullmatch(r"\d{8}-\d{6}_01_[a-z0-9-]+(_marco)?\.png", a.name) for a in archivos)
    assert sum(a.name.endswith("_marco.png") for a in archivos) == 1
    assert Path(r.captura).resolve() in [a.resolve() for a in archivos]
    assert not r.captura.endswith("_marco.png")
    tipos = [t for _, t in p.top.capturas]
    assert "pagina_completa" in tipos and f"elemento:{S('marco_no_molestar')}" in tipos


def test_enviar_solicitud_sin_verificar_guarda_diagnostico_tras_enviar(nuevo_portal, crear_bot, tmp_path, monkeypatch):
    p = _listo(nuevo_portal)
    capturas = tmp_path / "salida" / "capturas"
    llamadas: list[tuple[Path, int, int]] = []
    monkeypatch.setattr(nm, "SELECTORES_VERIFICADOS", False)
    monkeypatch.setattr(nm, "capturar_pagina", lambda page, d, n: llamadas.append((Path(d), n, p.envios)) or [])
    r = crear_bot(p.page, capturas_dir=capturas).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "ok"
    assert llamadas == [(capturas / "diagnostico", 1, 1)]  # una vez, después del POST


def test_enviar_solicitud_verificada_no_guarda_diagnostico(nuevo_portal, crear_bot, tmp_path, monkeypatch):
    p = _listo(nuevo_portal)
    llamadas: list[Any] = []
    monkeypatch.setattr(nm, "SELECTORES_VERIFICADOS", True)
    monkeypatch.setattr(nm, "capturar_pagina", lambda *a: llamadas.append(a) or [])
    r = crear_bot(p.page, capturas_dir=tmp_path / "capturas").enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "ok"
    assert llamadas == []


def test_enviar_solicitud_diagnostico_que_falla_no_cambia_el_resultado(nuevo_portal, crear_bot, tmp_path, monkeypatch):
    p = _listo(nuevo_portal)
    monkeypatch.setattr(nm, "SELECTORES_VERIFICADOS", False)

    def falla(*_a: Any) -> list[Path]:
        raise OSError("disco lleno")

    monkeypatch.setattr(nm, "capturar_pagina", falla)
    r = crear_bot(p.page, capturas_dir=tmp_path / "capturas").enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "ok"
    assert p.envios == 1


def test_enviar_solicitud_reingresa_una_vez_si_la_sesion_expiro(nuevo_portal, crear_bot):
    p = _en_login(nuevo_portal)
    bot = crear_bot(p.page)
    bot.iniciar_sesion(nm.Secreto(SECRETO_CU))
    p.expirar_en_goto_no_molestar = p.gotos_no_molestar + 1
    r = bot.enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    assert r.estado == "ok"
    assert p.envios_login == [("123456785", SECRETO_CU), ("123456785", SECRETO_CU)]
    assert p.envios == 1


def test_enviar_solicitud_captcha_antes_de_finalizar_pausa(nuevo_portal, crear_bot, capsys):
    p = _listo(nuevo_portal, captcha_al_enviar=True)
    preguntar = PreguntarGuionado([lambda prompt: (p.resolver_captcha(), "")[1]])
    r = crear_bot(p.page, preguntar=preguntar).enviar_solicitud(fila_real(1), ENTRADAS_CLARO)
    salida = capsys.readouterr()
    assert r.estado == "ok"
    assert preguntar.respuestas == []
    assert "captcha" in (preguntar.texto + salida.out + salida.err).lower()
    assert p.envios == 1


# =============================================================================
# main()
# =============================================================================
def _textos_de_salida(salida_dir: Path, capturado, tracker: Path) -> list[str]:
    textos = [capturado.out, capturado.err, todas_las_celdas(tracker)]
    textos += [p.read_text(encoding="utf-8", errors="replace") for p in salida_dir.rglob("*.log")]
    textos += [p.read_text(encoding="utf-8-sig", errors="replace") for p in salida_dir.rglob("*.csv")]
    return textos


def _cambiadas(antes: Path, despues: Path) -> set[tuple[str, str]]:
    a, d = celdas(antes), celdas(despues)
    return {k for k in a.keys() | d.keys() if a.get(k) != d.get(k)}


def test_main_dry_run_con_copias_reales_muestra_plan_y_no_escribe_nada(archivos, sin_efectos, capsys):
    antes_tracker, antes_mapeo = archivos.tracker.read_bytes(), archivos.mapeo.read_bytes()
    entrada_antes = sorted(p.name for p in archivos.dir.iterdir())
    codigo = nm.main(archivos.argv("--dry-run"))
    salida = capsys.readouterr()
    texto = salida.out + salida.err
    bajo = texto.lower()
    assert codigo == 0
    assert sin_efectos.llamadas == []
    assert archivos.tracker.read_bytes() == antes_tracker
    assert archivos.mapeo.read_bytes() == antes_mapeo
    assert sorted(p.name for p in archivos.dir.iterdir()) == entrada_antes
    for empresa in ("Claro Chile", "Entel", "Banco Security", "Isapre Consalud"):
        assert empresa in texto
    assert "pendiente con 2 entradas" in bajo
    assert "pendiente con 3 entradas" in bajo
    assert "sin entrada en sernac" in bajo
    assert "no aplica" in bajo
    previas = sorted((archivos.salida / "vista_previa").glob("*.txt"))
    assert len(previas) == 30
    textos = [p.read_text(encoding="utf-8") for p in previas]
    assert not any("{" in t or "}" in t for t in textos)
    entel = next(t for t in textos if "protecciondedatos@entel.cl" in t)
    assert "Para" in entel and "Asunto" in entel
    assert "Oposición al tratamiento de datos – Entel" in entel
    assert "Empresa Nacional de Telecomunicaciones S.A. / Entel PCS Telecomunicaciones S.A." in entel
    assert "92.580.000-7 / 96.806.980-2" in entel
    assert "12.345.678-5" in entel and "+56 9 1234 5678" in entel and "Ana Pérez Soto" in entel
    respaldos = archivos.salida / "respaldos"
    assert not respaldos.exists() or not any(respaldos.iterdir())


def test_main_dry_run_avisa_empresas_sin_mapeo(archivos, sin_efectos, capsys):
    lineas = archivos.mapeo.read_text(encoding="utf-8-sig").splitlines()
    archivos.mapeo.write_text("\n".join(ln for ln in lineas if not ln.startswith("50;")) + "\n", encoding="utf-8-sig")
    codigo = nm.main(archivos.argv("--dry-run"))
    assert codigo == 0
    assert "falta mapeo" in capsys.readouterr().out.lower()


def test_main_dry_run_marca_lo_ya_hecho(archivos, sin_efectos, capsys):
    wb = openpyxl.load_workbook(archivos.tracker)
    wb["Empresas"]["O2"] = datetime(2026, 10, 1)
    wb["Empresas"]["P2"] = "1111111"
    wb["Empresas"]["L2"] = datetime(2026, 10, 1)
    wb.save(archivos.tracker)
    codigo = nm.main(archivos.argv("--dry-run"))
    assert codigo == 0
    assert "ya hecho" in capsys.readouterr().out.lower()


def test_main_dry_run_con_rut_invalido_sale_con_2(archivos, sin_efectos, capsys):
    archivos.config.write_text(CONFIG_VALIDA.replace("12.345.678-5", "12.345.678-9"), encoding="utf-8")
    codigo = nm.main(archivos.argv("--dry-run"))
    salida = capsys.readouterr()
    assert codigo == 2
    assert "rut" in (salida.out + salida.err).lower()
    assert sin_efectos.llamadas == []


def test_main_dry_run_con_marcador_desconocido_sale_con_2_y_lo_nombra(archivos, sin_efectos, capsys):
    archivos.plantilla.write_text(PLANTILLA_VALIDA + "\nDato: {campo_inexistente}\n", encoding="utf-8")
    codigo = nm.main(archivos.argv("--dry-run"))
    salida = capsys.readouterr()
    assert codigo == 2
    assert "campo_inexistente" in salida.out + salida.err


def test_main_dry_run_con_tracker_sin_columna_de_salida_sale_con_2(archivos, sin_efectos, capsys):
    wb = openpyxl.load_workbook(archivos.tracker)
    wb["Empresas"]["P1"] = "Observaciones"
    wb.save(archivos.tracker)
    assert nm.main(archivos.argv("--dry-run")) == 2
    assert "N° solicitud SERNAC" in "".join(capsys.readouterr())


def test_main_dry_run_con_mapeo_mal_formado_sale_con_2(archivos, sin_efectos, capsys):
    archivos.mapeo.write_text("#;Marca;Entrada SERNAC\n1;Claro Chile;CLARO MOVIL\n", encoding="utf-8-sig")
    assert nm.main(archivos.argv("--dry-run")) == 2


def test_main_error_de_validacion_sale_con_2_antes_de_pedir_claves(archivos, sin_efectos, capsys):
    archivos.config.write_text(CONFIG_VALIDA.replace("+56 9 1234 5678", "123"), encoding="utf-8")
    assert nm.main(archivos.argv()) == 2
    assert sin_efectos.llamadas == []


def test_main_mapear_pide_solo_claveunica_actualiza_mapeo_y_no_toca_el_tracker(archivos, monkeypatch, tiempos_rapidos, capsys):
    lineas = archivos.mapeo.read_text(encoding="utf-8-sig").splitlines()
    archivos.mapeo.write_text("\n".join(ln for ln in lineas if not ln.startswith(("3;", "11;"))) + "\n", encoding="utf-8-sig")
    antes_tracker = archivos.tracker.read_bytes()
    BuscadorFalso.instancias.clear()
    pagina = PortalFalso().page
    lanzamientos = []

    @contextlib.contextmanager
    def falso_lanzar(*args, **kwargs):
        lanzamientos.append((args, kwargs))
        yield pagina

    monkeypatch.setattr(nm, "lanzar_navegador", falso_lanzar)
    monkeypatch.setattr(nm, "BotSernac", BuscadorFalso)
    monkeypatch.setattr(smtplib, "SMTP_SSL", _prohibido("smtplib.SMTP_SSL"))
    pedidos = []
    monkeypatch.setattr(getpass, "getpass", lambda prompt="", stream=None: pedidos.append(prompt) or SECRETO_CU)
    monkeypatch.setattr(builtins, "input", PreguntarGuionado(["", ""]))

    codigo = nm.main(archivos.argv("--mapear"))
    salida = capsys.readouterr()

    assert codigo == 0
    assert len(pedidos) == 1 and "clave" in pedidos[0].lower() and "gmail" not in pedidos[0].lower()
    assert len(lanzamientos) == 1
    assert archivos.tracker.read_bytes() == antes_tracker
    m = nm.Mapeo.cargar(archivos.mapeo)
    assert {e.texto for e in m.entradas(fila_de("x", numero=3))} == set(_textos("movistar"))
    assert m.entradas(fila_de("x", numero=11)) == [nm.EntradaSernac(None, "BANCO DE CHILE - BANCO EDWARDS - CITI")]
    assert m.entradas(fila_de("x", numero=1)) == ENTRADAS_CLARO
    assert m.entradas(fila_de("x", numero=18)) == []
    assert len(BuscadorFalso.instancias[-1].logins) == 1
    for texto in _textos_de_salida(archivos.salida, salida, archivos.tracker):
        assert SECRETO_CU not in texto


@pytest.fixture
def entorno_correo(monkeypatch, tiempos_rapidos):
    """main() en fase correo con SMTP falso; getpass devuelve la clave de aplicación con espacios."""
    from types import SimpleNamespace

    fab = FabricaSMTP()
    monkeypatch.setattr(smtplib, "SMTP_SSL", fab)
    pedidos: list[str] = []

    def falso_getpass(prompt="Password: ", stream=None):
        pedidos.append(prompt)
        return SECRETO_GMAIL

    monkeypatch.setattr(getpass, "getpass", falso_getpass)
    monkeypatch.setattr(nm, "lanzar_navegador", _prohibido("lanzar_navegador"))
    monkeypatch.setattr(builtins, "input", PreguntarGuionado())
    return SimpleNamespace(fab=fab, pedidos=pedidos, sleep=tiempos_rapidos)


def _filas_correo(path: Path) -> dict[int, str]:
    ws = openpyxl.load_workbook(path)["Empresas"]
    return {r: ws.cell(row=r, column=6).value for r in range(2, 52) if ws.cell(row=r, column=11).value == "Correo"}


def test_main_solo_correo_con_copia_real_envia_30_correos_y_solo_escribe_enviado(archivos, entorno_correo, capsys):
    pristino = archivos.dir / "pristino.xlsx"
    shutil.copy2(archivos.tracker, pristino)
    correo_por_fila = _filas_correo(pristino)
    codigo = nm.main(archivos.argv("--solo", "correo", "-v"))
    salida = capsys.readouterr()
    fab = entorno_correo.fab
    assert codigo == 0
    assert len(entorno_correo.pedidos) == 1 and "gmail" in entorno_correo.pedidos[0].lower()
    assert fab.logins[0] == (GMAIL_USUARIO, SECRETO_GMAIL_SIN_ESPACIOS)
    assert len(correo_por_fila) == 30
    assert sorted(str(m["To"]) for m in fab.enviados) == sorted(correo_por_fila.values())
    assert _cambiadas(pristino, archivos.tracker) == {("Empresas", f"L{r}") for r in correo_por_fila}
    despues = celdas(archivos.tracker)
    assert all(es_fecha(despues[("Empresas", f"L{r}")], date.today()) for r in correo_por_fila)
    assert entorno_correo.sleep.llamadas.count(70) == 29
    texto = salida.out + salida.err
    assert "Acción manual pendiente" in texto
    assert "https://legalmovistar.mvch.cl/" in texto and "WOM Chile" in texto
    assert len(list((archivos.salida / "respaldos").glob("*.xlsx"))) == 1
    resumenes = list(archivos.salida.glob("resumen_*.csv"))
    assert resumenes and ";" in resumenes[0].read_text(encoding="utf-8-sig").splitlines()[0]
    assert list(archivos.salida.rglob("*.log"))
    for t in _textos_de_salida(archivos.salida, salida, archivos.tracker):
        assert SECRETO_GMAIL not in t
        assert SECRETO_GMAIL_SIN_ESPACIOS not in t


def test_main_solo_correo_respeta_limite(archivos, entorno_correo, capsys):
    assert nm.main(archivos.argv("--solo", "correo", "--limite", "1")) == 0
    assert len(entorno_correo.fab.enviados) == 1


def test_main_solo_correo_con_auth_fallida_explica_y_no_envia(archivos, entorno_correo, capsys):
    fab = entorno_correo.fab
    fab.fallos["login"] = [smtplib.SMTPAuthenticationError(535, b"5.7.8 Username and Password not accepted")]
    antes = archivos.tracker.read_bytes()
    codigo = nm.main(archivos.argv("--solo", "correo"))
    salida = capsys.readouterr()
    texto = (salida.out + salida.err).lower()
    assert codigo == 1
    assert fab.enviados == []
    assert len(fab.logins) == 1
    assert "contrase" in texto and "aplicaci" in texto
    assert archivos.tracker.read_bytes() == antes


def test_main_clave_vacia_tres_veces_aborta_con_130(archivos, entorno_correo, monkeypatch, capsys):
    pedidos = []
    monkeypatch.setattr(getpass, "getpass", lambda prompt="", stream=None: pedidos.append(prompt) or "")
    assert nm.main(archivos.argv("--solo", "correo")) == 130
    assert 1 <= len(pedidos) <= 4
    assert entorno_correo.fab.conexiones == []


def test_main_ctrl_c_al_pedir_clave_sale_con_130(archivos, entorno_correo, monkeypatch, capsys):
    monkeypatch.setattr(getpass, "getpass", lanzar(KeyboardInterrupt()))
    assert nm.main(archivos.argv("--solo", "correo")) == 130
    assert entorno_correo.fab.conexiones == []


def test_main_solo_sernac_con_fakepage_registra_fechas_y_numeros(archivos, monkeypatch, tiempos_rapidos, capsys):
    portal = PortalFalso(resultados_envio=(("ok", "1234567"), ("ok", "7654321")))
    lanzamientos = []

    @contextlib.contextmanager
    def falso_lanzar(*args, **kwargs):
        lanzamientos.append((args, kwargs))
        yield portal.page

    monkeypatch.setattr(nm, "lanzar_navegador", falso_lanzar)
    monkeypatch.setattr(smtplib, "SMTP_SSL", _prohibido("smtplib.SMTP_SSL"))
    pedidos = []
    monkeypatch.setattr(getpass, "getpass", lambda prompt="", stream=None: pedidos.append(prompt) or SECRETO_CU)
    monkeypatch.setattr(builtins, "input", PreguntarGuionado())
    pristino = archivos.dir / "pristino.xlsx"
    shutil.copy2(archivos.tracker, pristino)

    codigo = nm.main(archivos.argv("--solo", "sernac", "--empresa", "entel", "--empresa", "banco de chile", "--sin-confirmar", "-v"))
    salida = capsys.readouterr()

    assert codigo == 0
    assert len(pedidos) == 1 and "clave" in pedidos[0].lower() and "gmail" not in pedidos[0].lower()
    assert len(lanzamientos) == 1
    assert not lanzamientos[0][1].get("headless", False) and True not in lanzamientos[0][0]
    assert nm.PORTAL_URL in portal.top.gotos
    assert portal.envios_login == [("123456785", SECRETO_CU)]
    assert [s["empresas"] for s in portal.solicitudes] == [
        [(e.id, e.texto) for e in ENTRADAS_ENTEL],
        [(27498, "BANCO DE CHILE - BANCO EDWARDS - CITI")],
    ]
    assert _cambiadas(pristino, archivos.tracker) == {("Empresas", c) for c in ("O5", "P5", "O12", "P12")}
    despues = celdas(archivos.tracker)
    assert es_fecha(despues[("Empresas", "O5")], date.today()) and es_fecha(despues[("Empresas", "O12")], date.today())
    assert (despues[("Empresas", "P5")], despues[("Empresas", "P12")]) == ("1234567", "7654321")
    assert len(list((archivos.salida / "capturas").glob("*.png"))) == 4
    assert tiempos_rapidos.llamadas.count(50) == 1
    for texto in _textos_de_salida(archivos.salida, salida, archivos.tracker):
        assert SECRETO_CU not in texto
        assert JWT_PORTAL not in texto
