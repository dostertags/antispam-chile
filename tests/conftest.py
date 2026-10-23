"""Fixtures y dobles de prueba para no_molestar.py (spec revisión 3).

Nada aquí toca la red: SMTP, getpass, navegador y pausas son falsos. Los fixtures NO llaman
funciones de no_molestar (solo usan constantes y dataclasses), para que una implementación
ausente aparezca como FALLO del test y no como error del fixture.

Los archivos reales del usuario (tracker y mapeo) solo se leen o se COPIAN a tmp_path; un fixture
de sesión verifica al final que no cambiaron.
"""
from __future__ import annotations

import base64
import hashlib
import html
import logging
import re
import shutil
import smtplib
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Callable

import openpyxl
import pytest
from playwright.sync_api import TimeoutError as PWTimeoutError

import no_molestar as nm

RAIZ = Path(nm.__file__).resolve().parent
# Tracker con el diseño real: la plantilla publicada (limpia, igual para todos); el de trabajo del
# usuario solo se vigila para que ninguna prueba lo toque.
TRACKER_REAL = RAIZ / nm.TRACKER_PLANTILLA
TRACKER_TRABAJO = RAIZ / "tracker_no_molestar_50_empresas.xlsx"
NOMBRE_TRACKER = TRACKER_TRABAJO.name  # nombre con que se copian (el que usa el script por defecto)
MAPEO_REAL = RAIZ / "mapeo_sernac.csv"

# ----------------------------------------------------------------------------- datos
SECRETO_CU = "Zq9#ClaveUnicaPrueba!"
SECRETO_GMAIL = "qzxw vbnm plok ijuh"
SECRETO_GMAIL_SIN_ESPACIOS = "qzxwvbnmplokijuh"
GMAIL_USUARIO = "ana.optout@gmail.com"
CORREO_PERSONA = "ana.perez@example.cl"
TEL_PERSONA_E164 = "+56912345678"
# Token FALSO con forma de JWT (encabezado.carga.firma), como el p= del portal. Se arma por partes
# para que los escáneres de secretos (gitleaks) no lo confundan con uno real.
JWT_PORTAL = ".".join(
    base64.urlsafe_b64encode(parte.encode()).decode().rstrip("=")
    for parte in ('{"alg":"HS256"}', '{"nombre":"Ana Perez","correo":"ana@example.cl"}')
) + ".Zx9-firmaFalsa_77"

CONFIG_VALIDA = """\
[persona]
nombre = "Ana Pérez Soto"
rut = "12.345.678-5"
telefono = "+56 9 1234 5678"
correo = "ana.perez@example.cl"

[gmail]
usuario = "ana.optout@gmail.com"
"""

PLANTILLA_VALIDA = """\
Asunto: Oposición al tratamiento de datos – {empresa}

Señores {empresa} ({razon_social}, RUT {rut_empresa}), sector {sector}:

Yo, {nombre}, RUT {rut}, me opongo al uso de mis datos con fines publicitarios y pido que no me
contacten al teléfono {telefono} ni al correo {correo} (art. 28B Ley 19.496).

Saluda atentamente,
{nombre}
{fecha}
"""

ENCABEZADOS_REALES = [
    "#", "Marca", "Sector", "Razón social", "RUT", "Correo datos personales", "Formulario web",
    "Opt-out propio", "Canal alternativo", "Fuente", "Vía de envío", "Enviado (fecha)",
    "Plazo respuesta correo", "¿Respondió?", "No Molestar ingresado", "N° solicitud SERNAC",
    "Plazo No Molestar", "¿Siguen llamando?", "Aviso incumplimiento", "Estado", "Notas",
]
FORMATO_FECHA_REAL = "dd\\-mm\\-yyyy"

# Filas copiadas de la hoja real "Empresas" (textos largos abreviados).
FILAS_REALES: dict[int, dict[str, Any]] = {
    1: {"#": 1, "Marca": "Claro Chile", "Sector": "Telecom",
        "Correo datos personales": "datos.personales@clarochile.cl",
        "Formulario web": "https://denuncias.americamovil.com/",
        "Opt-out propio": "Formulario 'No Molestar' en el sitio de Claro",
        "Canal alternativo": "Centros de Atención Claro VTR (presencial) / call center",
        "Fuente": "https://www.clarochile.cl/portal/cl/recursos_contenido/pdf/politica.pdf",
        "Vía de envío": "Correo"},
    3: {"#": 3, "Marca": "Movistar Chile (Telefónica Chile)", "Sector": "Telecom",
        "Formulario web": "https://legalmovistar.mvch.cl/",
        "Opt-out propio": "Formulario No Molestar: https://catalogo.movistar.cl/planes/formularionomolestar",
        "Canal alternativo": "Teléfono 103, sucursales Movistar, redes sociales",
        "Fuente": "https://ww2.movistar.cl/centro-de-transparencia/derechos.html",
        "Vía de envío": "Formulario web"},
    4: {"#": 4, "Marca": "Entel", "Sector": "Telecom",
        "Razón social": "Empresa Nacional de Telecomunicaciones S.A. / Entel PCS Telecomunicaciones S.A.",
        "RUT": "92.580.000-7 / 96.806.980-2", "Correo datos personales": "protecciondedatos@entel.cl",
        "Opt-out propio": "Oposición por cada canal, siguiendo instrucciones de cada mensaje",
        "Canal alternativo": "Mi Entel (web), App Entel, sucursales Entel, call center",
        "Fuente": "https://entelconviene.entel.cl/politica.pdf", "Vía de envío": "Correo"},
    5: {"#": 5, "Marca": "WOM Chile", "Sector": "Telecom",
        "Canal alternativo": "Centro de ayuda WOM: https://www.wom.cl/centro-de-ayuda/",
        "Fuente": "https://www.wom.cl/politica-privacidad/", "Vía de envío": "Canal alternativo"},
    11: {"#": 11, "Marca": "Banco de Chile", "Sector": "Banca", "Razón social": "Banco de Chile",
         "RUT": "97.004.000-5", "Correo datos personales": "protecciondatos@bancochile.cl",
         "Canal alternativo": "'Derechos ARCO' channel on bancochile.cl website",
         "Fuente": "https://sitiospublicos.bancochile.cl/personas/politica-de-privacidad",
         "Vía de envío": "Correo"},
    18: {"#": 18, "Marca": "Banco Security", "Sector": "Banca",
         "Canal alternativo": "Servicio Atención Clientes Security +56 2 2584 4060",
         "Fuente": "https://personas.bancosecurity.cl/informacion-al-cliente/politica-de-privacidad",
         "Vía de envío": "Canal alternativo"},
}

# ------------------------------------------------- directorio SERNAC (búsquedas reales, WALK.md)
RESULTADOS_WALK: dict[str, list[tuple[int, str]]] = {
    "entel": [
        (269830, "RESTAURANT HUENTELAUQUEN"),
        (671334, "CELULAR CENTEL - LA SERENA (REPARACIONES - CENTRO DE LLAMADOS)"),
        (672519, "ZENTELLA MOTOS"),
        (713407, "PARENTELAS"),
        (719516, "GRC DENTELITE"),
        (432, "CARRIER 123 ENTEL S.A. (LARGA DISTANCIA - INTERNET)"),
        (184, "ENTELPHONE (TELEFONIA FIJA - INTERNET ADSL- NGN- WILL- SUPER ADSL) - EPHONE - VENTAS INTERNET"),
        (1956650, "HUENTELAUQUEN"),
        (183, "ENTEL PCS MOVIL (CELULAR - INTERNET MOVIL) - VENTAS INTERNET"),
    ],
    "claro": [
        (410878, "CLINICA MIGUEL CLARO"),
        (762335, "OPTICAS V CLARO"),
        (770746, "AGUAS RIO CLARO"),
        (793963, "COMERCIAL PURO Y CLARO (ARTICULOS DE LINEA BLANCA)"),
        (837206, "CLAROALTO SPA - INGENIERIA TELECOMUNICACION Y CONSTRUCCION CLAROALTO SPA"),
        (880277, "INMOBILIARIA CLARO SOLAR 636"),
        (830241, "INMOBILIARIA RIO CLARO S.A."),
        (1550624, "EL FOGON DE RIO CLARO"),
        (1201223, "INMOBILIARIA CLARO DECOMBE"),
        (178, "CLARO MOVIL (CELULAR - INTERNET MOVIL -PRODUCTO 3 PLAY INALAMBRICO) - VENTAS INTERNET"),
        (1543191, "FORESTAL RIO CLARO"),
        (1623244, "PREUNIVERSITARIO MONTE CLARO"),
        (1914175, "CLARO DE LUNA COLLEGE"),
        (1961086, "INMOBILIARIA MONTECLARO"),
        (1970686, "SUR CLARO"),
        (297054, "PELUQUERIA CLAROSCURO - LILIAN CECILIA GONZALEZ LORR"),
        (2089306, "CLINICA RIO CLARO"),
        (2114981, "CENTRO VETERINARIO EL CLARON"),
        (146420, "CLARO FIJO (TELEFONIA FIJA - TV CABLE - INTERNET DOMICILIARIA - CABLE SATELITAL - 3 PLAY ALAMBRICO)"),
        (163496, "GESTION INMOBILIARIA BEZANILLA CLARO"),
        (488442, "ENMARCACIONES CLAROSCURO"),
    ],
    "falabella": [
        (159, "BANCO FALABELLA -TARJETA DE CREDITO VISA BANCO FALABELLA"),
        (160, "TARJETA CMR FALABELLA - CMR FALABELLA VISA"),
        (173, "SEGUROS FALABELLA GENERALES Y VIDA - VENTAS INTERNET"),
        (723331, "FALABELLA INMOBILIARIO"),
        (402, "DESPEGAR.COM CHILE - VIAJES FALABELLA"),
        (885848, "SGF (SERVICIOS Y GARANTIAS FALABELLA SPA)"),
        (182, "TIENDAS FALABELLA"),
        (146573, "BANCO FALABELLA CORREDORES DE SEGUROS LIMITADA (CONTRATADO SOLO A TRAVES DEL BANCO)"),
        (579802, "FALABELLA.COM - LINIO"),
    ],
    "banco de chile": [
        (27498, "BANCO DE CHILE - BANCO EDWARDS - CITI"),
        (57701, "AFUBACH (ASOCIACION DE FUNCIONARIOS DEL BANCO DE CHILE)"),
    ],
    "movistar": [
        (175, "MOVISTAR TELEFONIA FIJA (TV DIGITAL- INTERNET DOMICILIARIA)-CARRIER 181- 188 Y 120- VENTAS INTERNET"),
        (1192505, "MOVISTAR MOVIL (TELEFONIA MOVIL - INTERNET MOVIL)"),
    ],
}


def _armar_directorio() -> list[tuple[int, str]]:
    vistos: set[int] = set()
    salida = []
    for bloque in RESULTADOS_WALK.values():
        for id_, texto in bloque:
            if id_ not in vistos:
                vistos.add(id_)
                salida.append((id_, texto))
    return salida


DIRECTORIO_SERNAC: list[tuple[int, str]] = _armar_directorio()
TEXTO_POR_ID = dict(DIRECTORIO_SERNAC)

ENTRADAS_CLARO = [
    nm.EntradaSernac(178, TEXTO_POR_ID[178]),
    nm.EntradaSernac(146420, TEXTO_POR_ID[146420]),
]
ENTRADAS_ENTEL = [nm.EntradaSernac(i, TEXTO_POR_ID[i]) for i in (183, 184, 432)]
ENTRADA_BANCO_CHILE = nm.EntradaSernac(27498, TEXTO_POR_ID[27498])


def _norm(texto: object) -> str:
    """Normalización propia de los dobles (no usa nm.normalizar)."""
    if texto is None:
        return ""
    s = unicodedata.normalize("NFKD", str(texto))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def buscar_en_directorio(termino: str, directorio=DIRECTORIO_SERNAC, limite: int | None = None) -> list[tuple[int, str]]:
    """Búsqueda por subcadena (como el buscador AJAX del portal), en el orden del directorio."""
    t = _norm(termino)
    res = [(i, x) for i, x in directorio if t and t in _norm(x)]
    return res[:limite] if limite else res


def tiempos_prueba(**cambios: Any) -> dict[str, Any]:
    """TIEMPOS con rangos distinguibles y timeouts cortos (todas las pausas pasan por sleep falso)."""
    t = dict(nm.TIEMPOS)
    t.update(
        pausa_entre_empresas=(50.0, 60.0),
        pausa_entre_correos=(70.0, 80.0),
        pausa_entre_acciones=(0.001, 0.002),
        pausa_entre_busquedas=(0.003, 0.004),
        timeout_elemento_ms=150,
        timeout_login_s=0.5,
        timeout_confirmacion_ms=150,
        retardo_tecleo_ms=0,
        reintentos=3,
        backoff_base_s=4,
    )
    t.update(cambios)
    return t


def fila_de(empresa: str, **kwargs: Any) -> nm.Fila:
    """Construye nm.Fila (dataclass) con valores por defecto razonables."""
    numero = kwargs.pop("numero", 1)
    datos = dict(
        numero_fila=(numero + 1) if isinstance(numero, int) else 2,
        numero=numero,
        empresa=empresa,
        razon_social="",
        rut_empresa="",
        correo_empresa="",
        via="Correo",
        formulario_web="",
        canal_alternativo="",
        valores={"#": numero, "Marca": empresa},
    )
    datos.update(kwargs)
    return nm.Fila(**datos)


def fila_real(numero: int, **kwargs: Any) -> nm.Fila:
    """Fila equivalente a la fila real `numero` de FILAS_REALES (como la entregaría Tracker)."""
    r = FILAS_REALES[numero]
    valores = {h: r.get(h) for h in ENCABEZADOS_REALES}
    for h in ("Plazo respuesta correo", "Plazo No Molestar", "Estado"):
        valores[h] = ""
    datos = dict(
        numero_fila=numero + 1,
        numero=numero,
        empresa=r["Marca"],
        razon_social=r.get("Razón social") or "",
        rut_empresa=r.get("RUT") or "",
        correo_empresa=r.get("Correo datos personales") or "",
        via=r["Vía de envío"],
        formulario_web=r.get("Formulario web") or "",
        canal_alternativo=r.get("Canal alternativo") or "",
        valores=valores,
    )
    datos.update(kwargs)
    return nm.Fila(**datos)


# --------------------------------------------------------------------- tracker Excel
def crear_tracker_real(
    path: Path,
    filas: list[dict[str, Any]],
    *,
    encabezados: list[str] | None = None,
    fila_encabezado: int = 1,
    hoja: str = "Empresas",
    con_resumen: bool = True,
    fila_en_blanco: bool = True,
    formatos_fecha: bool = True,
    otras_hojas: bool = True,
) -> Path:
    """Crea un tracker con la estructura de la hoja real (21 columnas A–U, fórmulas en «Plazo
    respuesta correo», «Plazo No Molestar» y «Estado», bloque «Resumen» con COUNTIF, hojas
    «Cómo usar» y «Feriados»). Cada fila es un dict encabezado -> valor; "#" se numera solo."""
    encabezados = list(encabezados or ENCABEZADOS_REALES)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = hoja
    if otras_hojas:
        como = wb.create_sheet("Cómo usar", 0)
        como["A1"] = "Tracker: 50 empresas – oposición y supresión de datos + No Molestar"
        fer = wb.create_sheet("Feriados")
        fer.append(["Fecha", "Feriado"])
        fer.append([datetime(2026, 10, 31), "Día de las Iglesias Evangélicas y Protestantes"])
    if fila_encabezado > 1:
        ws.cell(row=1, column=1, value="Seguimiento No Molestar — 50 empresas")
    col = {h: i for i, h in enumerate(encabezados, 1)}
    for h, i in col.items():
        ws.cell(row=fila_encabezado, column=i, value=h)

    def letra(h: str) -> str:
        return openpyxl.utils.get_column_letter(col[h])

    for n, fila in enumerate(filas, 1):
        desconocidas = set(fila) - set(encabezados)
        assert not desconocidas, f"fila de prueba con columnas inexistentes: {desconocidas}"
        r = fila_encabezado + n
        valores = dict(fila)
        if "#" in col and "#" not in valores:
            valores["#"] = n
        if {"Enviado (fecha)", "Plazo respuesta correo"} <= set(col):
            L = letra("Enviado (fecha)")
            valores.setdefault("Plazo respuesta correo", f'=IF({L}{r}="","",WORKDAY({L}{r},2,Feriados!$A$2:$A$60))')
        if {"No Molestar ingresado", "Plazo No Molestar"} <= set(col):
            O = letra("No Molestar ingresado")
            valores.setdefault("Plazo No Molestar", f'=IF({O}{r}="","",WORKDAY({O}{r},7,Feriados!$A$2:$A$60))')
        if "Estado" in col:
            valores.setdefault("Estado", f'=IF(AND(L{r}="",O{r}=""),"Por enviar","En plazo")')
        for h, v in valores.items():
            ws.cell(row=r, column=col[h], value=v)
        if formatos_fecha:
            for h in ("Enviado (fecha)", "No Molestar ingresado"):
                if h in col:
                    ws.cell(row=r, column=col[h]).number_format = FORMATO_FECHA_REAL
    ultima = fila_encabezado + len(filas)
    if con_resumen:
        inicio = ultima + (2 if fila_en_blanco else 1)
        ws.cell(row=inicio, column=2, value="Resumen")
        for k, (etiqueta, formula) in enumerate(
            [
                ("Por enviar", f'=COUNTIF(T2:T{ultima},"Por enviar")'),
                ("En plazo", f'=COUNTIF(T2:T{ultima},"En plazo")'),
                ("Con correo oficial", f'=COUNTIF(K2:K{ultima},"Correo")'),
                ("Solo canal alternativo", f'=COUNTIF(K2:K{ultima},"Canal alternativo")'),
            ],
            1,
        ):
            ws.cell(row=inicio + k, column=3, value=etiqueta)
            ws.cell(row=inicio + k, column=4, value=formula)
    wb.save(path)
    return path


def filas_reales(*numeros: int) -> list[dict[str, Any]]:
    return [dict(FILAS_REALES[n]) for n in numeros]


def encabezados_de(path: Path, fila_encabezado: int = 1, hoja: str = "Empresas") -> list[Any]:
    ws = openpyxl.load_workbook(path)[hoja]
    valores = [c.value for c in ws[fila_encabezado]]
    while valores and valores[-1] in (None, ""):
        valores.pop()
    return valores


def celda(path: Path, numero_fila: int, encabezado: str, *, fila_encabezado: int = 1, hoja: str = "Empresas") -> Any:
    """Celda openpyxl bajo `encabezado` en la fila de Excel indicada (fórmulas como texto)."""
    ws = openpyxl.load_workbook(path)[hoja]
    for col, c in enumerate(ws[fila_encabezado], 1):
        if c.value == encabezado:
            return ws.cell(row=numero_fila, column=col)
    raise AssertionError(f"no existe la columna {encabezado!r}; hay: {encabezados_de(path, fila_encabezado, hoja)}")


def valor_celda(path: Path, numero_fila: int, encabezado: str, **kwargs: Any) -> Any:
    return celda(path, numero_fila, encabezado, **kwargs).value


def celdas(path: Path) -> dict[tuple[str, str], Any]:
    """(hoja, coordenada) -> valor de todas las celdas no vacías del libro."""
    wb = openpyxl.load_workbook(path)
    salida = {}
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for c in fila:
                if c.value is not None:
                    salida[(ws.title, c.coordinate)] = c.value
    return salida


def todas_las_celdas(path: Path) -> str:
    return "\n".join(str(v) for v in celdas(path).values())


def es_fecha(valor: Any, dia: date) -> bool:
    """Acepta date/datetime de ese día."""
    if isinstance(valor, datetime):
        return valor.date() == dia
    if isinstance(valor, date):
        return valor == dia
    return False


def vacio(valor: Any) -> bool:
    return valor is None or valor == ""


# ----------------------------------------------------------------------- dobles simples
class PreguntarGuionado:
    """Doble de `preguntar`: devuelve respuestas guionadas en orden.

    Cada respuesta puede ser: str; una excepción (instancia o clase) que se lanza; o un callable
    f(prompt) -> str (simula que la persona hace algo en el navegador y luego responde).
    Si se agota, el test falla: la implementación preguntó algo no previsto.
    """

    def __init__(self, respuestas: list[Any] | tuple[Any, ...] = ()) -> None:
        self.respuestas = list(respuestas)
        self.prompts: list[str] = []

    def __call__(self, prompt: str = "") -> str:
        self.prompts.append(str(prompt))
        if not self.respuestas:
            pytest.fail(f"preguntar() llamado sin respuesta guionada; prompt={prompt!r}", pytrace=False)
        r = self.respuestas.pop(0)
        if isinstance(r, BaseException):
            raise r
        if isinstance(r, type) and issubclass(r, BaseException):
            raise r()
        if callable(r):
            return r(prompt)
        return r

    @property
    def texto(self) -> str:
        return "\n".join(self.prompts)


class SleepGrabador:
    """Doble de sleep: registra los segundos pedidos y no espera."""

    def __init__(self) -> None:
        self.llamadas: list[float] = []

    def __call__(self, segundos: float) -> None:
        self.llamadas.append(segundos)

    def en_rango(self, minimo: float, maximo: float) -> list[float]:
        return [s for s in self.llamadas if minimo <= s <= maximo]


class FabricaSMTP:
    """Reemplazo de smtplib.SMTP_SSL. Registra conexiones/login/envíos.

    `fallos` = {"conectar"|"login"|"send_message": [excepción o None, ...]}: se consume una por
    llamada (None = esa llamada funciona). Tras SMTPServerDisconnected la conexión queda muerta.
    """

    def __init__(self, fallos: dict[str, list[BaseException | None]] | None = None) -> None:
        self.fallos = {k: list(v) for k, v in (fallos or {}).items()}
        self.conexiones: list[tuple[Any, Any, dict[str, Any]]] = []
        self.logins: list[tuple[str, str]] = []
        self.enviados: list[Any] = []
        self.intentos_envio = 0
        self.quits = 0

    def _quizas_fallar(self, operacion: str) -> None:
        cola = self.fallos.get(operacion)
        if cola:
            exc = cola.pop(0)
            if exc is not None:
                raise exc

    def __call__(self, host: Any = None, port: Any = 0, *args: Any, **kwargs: Any) -> "FakeSMTP":
        self.conexiones.append((host, port, kwargs))
        self._quizas_fallar("conectar")
        return FakeSMTP(self)


class FakeSMTP:
    def __init__(self, fabrica: FabricaSMTP) -> None:
        self._fabrica = fabrica
        self.muerta = False

    def __enter__(self) -> "FakeSMTP":
        return self

    def __exit__(self, *exc: Any) -> bool:
        self.quit()
        return False

    def ehlo(self, *args: Any, **kwargs: Any) -> tuple[int, bytes]:
        return (250, b"ok")

    def noop(self) -> tuple[int, bytes]:
        if self.muerta:
            raise smtplib.SMTPServerDisconnected("conexión cerrada")
        return (250, b"ok")

    def login(self, user: str, password: str, *args: Any, **kwargs: Any) -> tuple[int, bytes]:
        if self.muerta:
            raise smtplib.SMTPServerDisconnected("conexión cerrada")
        self._fabrica.logins.append((user, password))
        self._fabrica._quizas_fallar("login")
        return (235, b"2.7.0 Accepted")

    def send_message(self, msg: Any, from_addr: Any = None, to_addrs: Any = None, *args: Any, **kwargs: Any) -> dict:
        if self.muerta:
            raise smtplib.SMTPServerDisconnected("conexión cerrada")
        self._fabrica.intentos_envio += 1
        try:
            self._fabrica._quizas_fallar("send_message")
        except smtplib.SMTPServerDisconnected:
            self.muerta = True
            raise
        self._fabrica.enviados.append(msg)
        return {}

    def quit(self) -> tuple[int, bytes]:
        self._fabrica.quits += 1
        self.muerta = True
        return (221, b"bye")

    def close(self) -> None:
        self.muerta = True


# ------------------------------------------------------------ DOM y página Playwright falsos
BBOX_VISIBLE = {"x": 10.0, "y": 10.0, "width": 120.0, "height": 30.0}
BBOX_NULO = {"x": 0.0, "y": 0.0, "width": 0.0, "height": 0.0}
CAJA_BADGE = {"x": 1014.0, "y": 826.0, "width": 256.0, "height": 60.0}  # medida real del badge v3
CAJA_V2 = {"x": 480.0, "y": 300.0, "width": 304.0, "height": 78.0}
CAJA_MARCO = {"x": 90.0, "y": 380.0, "width": 1100.0, "height": 900.0}
PNG_FALSO = b"\x89PNG\r\n\x1a\nPNG-FALSO"

SUPERFICIE_PAGINA = (
    "goto url title content screenshot frames frame_locator wait_for_timeout wait_for_load_state "
    "evaluate locator"
)
SUPERFICIE_MARCO = "locator"
SUPERFICIE_LOCATOR = (
    "first nth count is_visible is_enabled is_checked check fill click press_sequentially "
    "select_option input_value inner_text all_inner_texts wait_for bounding_box screenshot"
)

URL_HOME = "https://www.sernac.cl/app/consumidor/index.php?c=home&a=home"
URL_CU = "https://accounts.claveunica.gob.cl/accounts/login/?next=/openid/authorize%3Fclient_id%3Dsernac"
URL_MARCO_LISTADO = f"https://www.sernac.cl/no-molestar/solicitudes?p={JWT_PORTAL}"
URL_MARCO_NUEVO = f"https://www.sernac.cl/no-molestar/solicitudes/new?p={JWT_PORTAL}"
URL_BADGE_INVISIBLE = (
    "https://www.google.com/recaptcha/enterprise/anchor?ar=1&k=6LdVkKAaPRUEBA&co=aHR0cHM6&hl=es"
    "&v=pCoGBhjs9s8EhFOHJFe8cqis&size=invisible&cb=7nq0x"
)
URL_ANCHOR_V2 = "https://www.google.com/recaptcha/api2/anchor?ar=1&k=6LcPrueba&co=aHR0&hl=es&size=normal"
URL_BFRAME = "https://www.google.com/recaptcha/api2/bframe?hl=es&v=abc&k=6LcPrueba"
URL_HCAPTCHA = "https://newassets.hcaptcha.com/captcha/v1/abc/static/hcaptcha.html#frame=checkbox"
URL_TURNSTILE = "https://challenges.cloudflare.com/cdn-cgi/challenge-platform/h/b/turnstile/if/ov2/av0/abc"

# Selectores que en un navegador real SÍ encuentran el iframe invisible del reCAPTCHA v3
# (256x60, visible). Una implementación que los use daría un falso positivo.
SELECTORES_QUE_VEN_EL_BADGE = (
    ".grecaptcha-badge",
    ".grecaptcha-badge iframe",
    "iframe[title='reCAPTCHA']",
    "iframe[src*='recaptcha']",
    "iframe[src*='recaptcha'][src*='anchor']",
    "iframe[src*='anchor']",
    "iframe[src*='size=invisible']",
    "iframe[src*='google.com/recaptcha']",
)


class ErrorModoEstricto(Exception):
    """Como el «strict mode violation» de Playwright: acción sobre un locator con >1 elemento.
    No es un error de Playwright a propósito: un bug así no debe esconderse tras reintentos."""


@dataclass
class Elemento:
    texto: str = ""
    visible: bool = True
    habilitado: bool = True
    marcado: bool = False
    valor: str = ""
    etiqueta: str = ""
    opciones: list[tuple[str, str]] = field(default_factory=list)  # (value, label) de un <select>
    caja: dict[str, float] | None = None
    datos: dict[str, Any] = field(default_factory=dict)
    al_click: Callable[..., Any] | None = None  # f(dom, elemento)
    al_llenar: Callable[..., Any] | None = None  # f(dom, elemento, valor)
    al_escribir: Callable[..., Any] | None = None  # f(dom, elemento, texto)
    al_marcar: Callable[..., Any] | None = None  # f(dom, elemento)
    al_seleccionar: Callable[..., Any] | None = None  # f(dom, elemento)
    clicks: int = 0
    llenados: list[str] = field(default_factory=list)
    escrito: list[str] = field(default_factory=list)
    checks: int = 0
    selecciones: list[str] = field(default_factory=list)


class DOMFalso:
    """Estado de un documento falso: selector (string EXACTO) -> lista de elementos.
    Lo manipulan los tests y PortalFalso; la implementación solo lo ve a través de FakePage."""

    def __init__(self, url: str = "about:blank", titulo: str = "", texto: str = "") -> None:
        self.url = url
        self.titulo = titulo
        self.texto = texto
        self.elementos: dict[str, list[Elemento]] = {}
        self.frames: list[str] = []
        self.status = 200
        self.gotos: list[str] = []
        self.capturas: list[tuple[str, str]] = []
        self.esperas: list[float] = []
        self.evaluaciones: list[str] = []
        self.consultados: list[str] = []
        self.al_goto: Callable[..., Any] | None = None
        self.al_evaluar: Callable[..., Any] | None = None

    def reiniciar(self, url: str | None = None, titulo: str | None = None, texto: str = "") -> None:
        """Nuevo documento (navegación): borra elementos y frames; conserva el historial."""
        self.elementos = {}
        self.frames = []
        self.texto = texto
        if url is not None:
            self.url = url
        if titulo is not None:
            self.titulo = titulo

    def agregar(self, selector: str, **kwargs: Any) -> Elemento:
        el = Elemento(**kwargs)
        self.elementos.setdefault(selector, []).append(el)
        return el

    def poner(self, selector: str, elementos: list[Elemento]) -> None:
        self.elementos[selector] = list(elementos)

    def quitar(self, *selectores: str) -> None:
        for s in selectores:
            self.elementos.pop(s, None)

    def todos(self, selector: str) -> list[Elemento]:
        return self.elementos.get(selector, [])

    def el(self, selector: str, i: int = 0) -> Elemento:
        return self.elementos[selector][i]

    def textos(self, selector: str) -> list[str]:
        return [e.texto for e in self.todos(selector)]

    def mostrar(self, *selectores: str) -> None:
        for s in selectores:
            for e in self.elementos.get(s, []):
                e.visible = True

    def ocultar(self, *selectores: str) -> None:
        for s in selectores:
            for e in self.elementos.get(s, []):
                e.visible = False

    def texto_visible(self) -> str:
        partes = [self.texto] if self.texto else []
        for lista in self.elementos.values():
            partes.extend(e.texto for e in lista if e.visible and e.texto)
        return "\n".join(partes)


class RespuestaFalsa:
    def __init__(self, status: int, url: str) -> None:
        self.status = status
        self.url = url

    @property
    def ok(self) -> bool:
        return 200 <= self.status < 400


class MarcoFalso:
    """Elemento de page.frames: solo .url y .name."""

    def __init__(self, url: str) -> None:
        self.url = url
        self.name = ""

    def __getattr__(self, nombre: str) -> Any:
        raise AttributeError(f"MarcoFalso solo implementa .url/.name (pedido: {nombre!r})")


def _escribir_png(path: Any) -> bytes:
    if path is not None:
        destino = Path(path)
        destino.parent.mkdir(parents=True, exist_ok=True)
        destino.write_bytes(PNG_FALSO)
    return PNG_FALSO


class FakeLocator:
    """Subconjunto de playwright Locator (spec §11). Busca el selector por string EXACTO."""

    def __init__(self, dom: DOMFalso, selector: str, indice: int | None = None, registro: DOMFalso | None = None) -> None:
        self._dom = dom
        self._selector = selector
        self._indice = indice
        self._registro = registro or dom  # dónde se anotan las capturas

    def __getattr__(self, nombre: str) -> Any:
        raise AttributeError(
            f"FakeLocator no implementa {nombre!r}; superficie permitida (spec §11): {SUPERFICIE_LOCATOR}"
        )

    def __repr__(self) -> str:
        return f"FakeLocator({self._selector!r}, indice={self._indice})"

    # -- internos
    def _todos(self) -> list[Elemento]:
        if self._selector in ("body", "html", ":root"):
            return [Elemento(texto=self._dom.texto_visible(), visible=True)]
        return self._dom.todos(self._selector)

    def _uno(self, accion: str) -> Elemento | None:
        els = self._todos()
        if self._indice is None:
            if len(els) > 1:
                raise ErrorModoEstricto(
                    f"strict mode violation: {self._selector!r} resolvió a {len(els)} elementos ({accion}); usa .first/.nth()"
                )
            return els[0] if els else None
        return els[self._indice] if 0 <= self._indice < len(els) else None

    def _accionable(self, accion: str) -> Elemento:
        el = self._uno(accion)
        if el is None or not el.visible:
            raise PWTimeoutError(f"Timeout: {self._selector!r} no visible ({accion})")
        if not el.habilitado:
            raise PWTimeoutError(f"Timeout: {self._selector!r} no está habilitado ({accion})")
        return el

    def _existente(self, accion: str) -> Elemento:
        el = self._uno(accion)
        if el is None:
            raise PWTimeoutError(f"Timeout: {self._selector!r} no existe ({accion})")
        return el

    # -- superficie Playwright
    @property
    def first(self) -> "FakeLocator":
        return FakeLocator(self._dom, self._selector, 0 if self._indice is None else self._indice, self._registro)

    def nth(self, indice: int) -> "FakeLocator":
        if self._indice is not None:
            return FakeLocator(self._dom, self._selector, self._indice if indice == 0 else -1, self._registro)
        return FakeLocator(self._dom, self._selector, indice, self._registro)

    def count(self) -> int:
        els = self._todos()
        if self._indice is None:
            return len(els)
        return 1 if 0 <= self._indice < len(els) else 0

    def is_visible(self, **kwargs: Any) -> bool:
        el = self._uno("is_visible")
        return bool(el is not None and el.visible)

    def is_enabled(self, **kwargs: Any) -> bool:
        return self._existente("is_enabled").habilitado

    def is_checked(self, **kwargs: Any) -> bool:
        return self._existente("is_checked").marcado

    def check(self, **kwargs: Any) -> None:
        el = self._accionable("check")
        if el.marcado:
            return
        if el.datos.get("intercepta") and not kwargs.get("force"):
            # radio/casilla «custom-control» de Bootstrap 4: el <input> queda bajo su <label>
            raise PWTimeoutError(
                f"Locator.check: Timeout exceeded.\n  - <label for=...> intercepts pointer events ({self._selector!r})"
            )
        el.marcado = True
        el.checks += 1
        if el.al_marcar:
            el.al_marcar(self._dom, el)

    def fill(self, value: str, **kwargs: Any) -> None:
        el = self._accionable("fill")
        el.valor = value
        el.llenados.append(value)
        if el.al_llenar:
            el.al_llenar(self._dom, el, value)

    def click(self, **kwargs: Any) -> None:
        el = self._accionable("click")
        el.clicks += 1
        if el.al_click:
            el.al_click(self._dom, el)

    def press_sequentially(self, text: str, **kwargs: Any) -> None:
        el = self._accionable("press_sequentially")
        el.valor += text
        el.escrito.append(text)
        if el.al_escribir:
            el.al_escribir(self._dom, el, text)

    def select_option(self, value: Any = None, *, index: Any = None, label: Any = None, **kwargs: Any) -> list[str]:
        el = self._accionable("select_option")
        elegida = None
        for i, (v, lab) in enumerate(el.opciones):
            if (label is not None and lab == label) or (value is not None and value in (v, lab)) or index == i:
                elegida = (v, lab)
                break
        if elegida is None:
            raise PWTimeoutError(f"Timeout: {self._selector!r} no tiene la opción value={value!r} label={label!r} index={index!r}")
        el.valor, el.etiqueta = elegida
        el.selecciones.append(elegida[1])
        if el.al_seleccionar:
            el.al_seleccionar(self._dom, el)
        return [elegida[0]]

    def input_value(self, **kwargs: Any) -> str:
        return self._existente("input_value").valor

    def inner_text(self, **kwargs: Any) -> str:
        return self._existente("inner_text").texto

    def all_inner_texts(self) -> list[str]:
        if self._indice is not None:
            el = self._uno("all_inner_texts")
            return [el.texto] if el else []
        return [e.texto for e in self._todos()]

    def wait_for(self, state: str = "visible", timeout: float | None = None, **kwargs: Any) -> None:
        el = self._uno("wait_for")
        ok = {
            "visible": el is not None and el.visible,
            "attached": el is not None,
            "hidden": el is None or not el.visible,
            "detached": el is None,
        }.get(state)
        if ok is None:
            raise ValueError(f"estado desconocido para wait_for: {state!r}")
        if not ok:
            raise PWTimeoutError(f"Timeout {timeout}ms esperando {self._selector!r} en estado {state!r}")

    def bounding_box(self, **kwargs: Any) -> dict[str, float] | None:
        el = self._uno("bounding_box")
        if el is None or not el.visible:
            return None
        return dict(el.caja if el.caja is not None else BBOX_VISIBLE)

    def screenshot(self, path: Any = None, **kwargs: Any) -> bytes:
        self._accionable("screenshot")
        datos = _escribir_png(path)
        self._registro.capturas.append((str(path or ""), f"elemento:{self._selector}"))
        return datos


class FakeFrameLocator:
    """Subconjunto de playwright FrameLocator: solo .locator(). Resuelve el documento del iframe
    en cada llamada (como Playwright), con su PROPIO mapa selector -> elementos."""

    def __init__(self, pagina: "FakePage", selector: str) -> None:
        self._pagina = pagina
        self._selector = selector

    def __getattr__(self, nombre: str) -> Any:
        raise AttributeError(
            f"FakeFrameLocator no implementa {nombre!r}; superficie permitida (spec §11): {SUPERFICIE_MARCO}"
        )

    def locator(self, selector: str) -> FakeLocator:
        dom = self._pagina._marcos.get(self._selector) or DOMFalso()
        dom.consultados.append(selector)
        return FakeLocator(dom, selector, registro=self._pagina._dom)


class FakePage:
    """Subconjunto de playwright Page (spec §11). Cualquier otro atributo -> AttributeError."""

    def __init__(self, dom: DOMFalso, marcos: dict[str, DOMFalso] | None = None) -> None:
        self._dom = dom
        self._marcos = dict(marcos or {})

    def __getattr__(self, nombre: str) -> Any:
        raise AttributeError(
            f"FakePage no implementa {nombre!r}; superficie permitida (spec §11): {SUPERFICIE_PAGINA}"
        )

    @property
    def url(self) -> str:
        return self._dom.url

    @property
    def frames(self) -> list[MarcoFalso]:
        return [MarcoFalso(self._dom.url)] + [MarcoFalso(u) for u in self._dom.frames]

    def goto(self, url: str, **kwargs: Any) -> RespuestaFalsa:
        self._dom.gotos.append(url)
        self._dom.url = url
        status = self._dom.status
        if self._dom.al_goto:
            devuelto = self._dom.al_goto(self._dom, url)
            if isinstance(devuelto, int):
                status = devuelto
        return RespuestaFalsa(status, url)

    def title(self) -> str:
        return self._dom.titulo

    def content(self) -> str:
        titulo = html.escape(self._dom.titulo)
        cuerpo = html.escape(self._dom.texto_visible())
        return f"<html><head><title>{titulo}</title></head><body>{cuerpo}</body></html>"

    def screenshot(self, path: Any = None, full_page: bool = False, **kwargs: Any) -> bytes:
        datos = _escribir_png(path)
        self._dom.capturas.append((str(path or ""), "pagina_completa" if full_page else "pagina"))
        return datos

    def wait_for_timeout(self, timeout: float) -> None:
        self._dom.esperas.append(timeout)

    def wait_for_load_state(self, state: str | None = None, **kwargs: Any) -> None:
        return None

    def evaluate(self, expression: str, arg: Any = None) -> Any:
        self._dom.evaluaciones.append(expression)
        if self._dom.al_evaluar:
            return self._dom.al_evaluar(expression, arg)
        if "innerText" in expression or "textContent" in expression:
            return self._dom.texto_visible()
        return None

    def locator(self, selector: str) -> FakeLocator:
        self._dom.consultados.append(selector)
        return FakeLocator(self._dom, selector)

    def frame_locator(self, selector: str) -> FakeFrameLocator:
        self._dom.consultados.append(f"frame_locator:{selector}")
        return FakeFrameLocator(self, selector)


def S(clave: str, i: int = 0) -> str:
    """Selector i de nm.SELECTORES[clave] (el portal falso registra sus elementos bajo estos)."""
    return nm.SELECTORES[clave][i]


def marcar_captcha(dom: DOMFalso, contiene: tuple[str, ...], *, visible: bool = True, caja: dict[str, float] | None = None) -> list[str]:
    """Registra como elementos los nm.SELECTORES_CAPTCHA que contienen alguno de los textos dados."""
    usados = []
    for s in nm.SELECTORES_CAPTCHA:
        if any(c in s for c in contiene):
            dom.poner(s, [Elemento(visible=visible, caja=caja if caja is not None else CAJA_V2)])
            usados.append(s)
    return usados


def quitar_captchas(dom: DOMFalso) -> None:
    for s in nm.SELECTORES_CAPTCHA:
        dom.quitar(s)
    dom.frames = [u for u in dom.frames if u not in (URL_ANCHOR_V2, URL_BFRAME, URL_HCAPTCHA, URL_TURNSTILE)]


def poner_badge_invisible(dom: DOMFalso) -> None:
    """reCAPTCHA v3/Enterprise como en el portal real: iframe size=invisible de 256x60 VISIBLE."""
    for s in SELECTORES_QUE_VEN_EL_BADGE:
        dom.poner(s, [Elemento(visible=True, caja=dict(CAJA_BADGE))])
    if URL_BADGE_INVISIBLE not in dom.frames:
        dom.frames.append(URL_BADGE_INVISIBLE)


# ---------------------------------------------------------------- portal falso (estructura real)
class PortalFalso:
    """Portal del Consumidor + ClaveÚnica + iframe «No Molestar» falsos, con la estructura grabada
    en el portal real (WALK.md). Selectores = primera alternativa de nm.SELECTORES. Los elementos
    del iframe viven en un DOM aparte, accesible solo con page.frame_locator("#frameNoMolestar").

    login: "ok" | "error" | "2fa" | "nada" | "captcha" (aparece un CAPTCHA al ingresar).
    resultados_envio: por cada clic en #finalizar: ("ok", N°) | ("sin_numero",) | ("listado",) |
    ("folio_en_pagina", folio) | ("nada",) | ("error_click",); el último se repite.
    """

    def __init__(
        self,
        *,
        login: str = "ok",
        directorio: list[tuple[int, str]] | None = None,
        limite_resultados: int | None = None,
        telefonos_precargados: tuple[str, ...] = (TEL_PERSONA_E164,),
        correos_precargados: tuple[str, ...] = (),
        acepto_marcado: bool = True,
        resultados_envio: tuple[tuple[Any, ...], ...] = (("ok", "1234567"),),
        status_portal: int = 200,
        canales_tel_desmarcados: tuple[int, ...] = (),
        canales_mail_desmarcados: tuple[int, ...] = (),
        rechazar_telefono: bool = False,
        captcha_en_login: bool = False,
        captcha_al_enviar: bool = False,
        modal_sesion: bool = False,
    ) -> None:
        self.login = login
        self.directorio = list(directorio if directorio is not None else DIRECTORIO_SERNAC)
        self.limite_resultados = limite_resultados
        self.telefonos_precargados = tuple(telefonos_precargados)
        self.correos_precargados = tuple(correos_precargados)
        self.acepto_marcado = acepto_marcado
        self.resultados_envio = tuple(resultados_envio)
        self.status_portal = status_portal
        self.canales_tel_desmarcados = set(canales_tel_desmarcados)
        self.canales_mail_desmarcados = set(canales_mail_desmarcados)
        self.rechazar_telefono = rechazar_telefono
        self.captcha_en_login = captcha_en_login
        self.captcha_al_enviar = captcha_al_enviar
        self.modal_sesion = modal_sesion
        # estado observable por los tests
        self.logueado = False
        self.envios = 0
        self.solicitudes: list[dict[str, Any]] = []
        self.busquedas: list[str] = []
        self.eventos: list[Any] = []
        self.envios_login: list[tuple[str, str]] = []
        self.contactos_agregados: list[tuple[str, str, str | None]] = []
        self.empresas_agregadas: list[tuple[int, str]] = []
        self.clicks_nueva = 0
        self.clicks_continuar = 0
        self.clicks_mantener = 0
        self.gotos_no_molestar = 0
        self.fallos_nueva_solicitud: list[BaseException] = []
        self.expirar_en_goto_no_molestar: int | None = None
        self.top_cu: dict[str, Elemento] | None = None  # campos de la última pantalla de ClaveÚnica
        self._borrador: dict[str, Any] = {}
        self.top = DOMFalso()
        self.marco = DOMFalso()
        self.top.al_goto = self._goto
        self.page = FakePage(self.top, {S("marco_no_molestar"): self.marco})

    # ------------------------------------------------------------- navegación
    def _goto(self, dom: DOMFalso, url: str) -> int:
        self.eventos.append(("goto", url))
        if url == nm.PORTAL_URL:
            if self.status_portal >= 400:
                self.top.reiniciar(url=url, titulo="403 Forbidden", texto="403 Forbidden nginx")
                self.marco.reiniciar()
                return self.status_portal
            self.home() if self.logueado else self.pantalla_login()
            return 200
        if url == nm.URL_NO_MOLESTAR:
            self.gotos_no_molestar += 1
            if self.expirar_en_goto_no_molestar == self.gotos_no_molestar:
                self.logueado = False
            self.pantalla_no_molestar() if self.logueado else self.pantalla_login()
            return 200
        self.top.reiniciar(url=url, titulo="No encontrado", texto="404")
        self.marco.reiniciar()
        return 404

    def pantalla_login(self) -> None:
        self.top.reiniciar(url=nm.PORTAL_URL, titulo="Portal del Consumidor", texto="Ingresa al Portal del Consumidor")
        self.marco.reiniciar()
        self.top.agregar(S("login_boton_claveunica"), texto="ClaveÚnica", al_click=lambda d, e: self.pantalla_cu())
        self.top.agregar("#btnRegistrado", texto="Clave SERNAC")
        poner_badge_invisible(self.top)

    def pantalla_cu(self) -> None:
        self.top.reiniciar(url=URL_CU, titulo="ClaveÚnica", texto="Ingresa tu RUN y tu ClaveÚnica")
        self.top_cu = {
            "run": self.top.agregar(S("cu_run"), al_escribir=self._tecla),
            "clave": self.top.agregar(S("cu_clave"), al_escribir=self._tecla),
            "ingresar": self.top.agregar(S("cu_ingresar"), texto="INGRESA", habilitado=False, al_click=self._ingresar),
        }
        poner_badge_invisible(self.top)
        if self.captcha_en_login:
            self.mostrar_captcha_v2()

    def _tecla(self, dom: DOMFalso, el: Elemento, texto: str) -> None:
        """keyup real: solo press_sequentially lo dispara (fill no)."""
        el.datos["keyup"] = True
        run, clave = dom.el(S("cu_run")), dom.el(S("cu_clave"))
        if run.datos.get("keyup") and clave.datos.get("keyup") and run.valor and clave.valor:
            dom.el(S("cu_ingresar")).habilitado = True

    def _ingresar(self, dom: DOMFalso, el: Elemento) -> None:
        self.envios_login.append((dom.el(S("cu_run")).valor, dom.el(S("cu_clave")).valor))
        if self.hay_captcha_visible():
            return
        if self.login == "ok":
            self.entrar()
        elif self.login == "error":
            dom.agregar(S("cu_error"), texto="RUN o ClaveÚnica incorrectos")
        elif self.login == "2fa":
            dom.agregar(nm.SELECTORES_2FA[0])
            dom.texto = "Ingresa el código de verificación enviado a tu celular"
        elif self.login == "captcha":
            self.mostrar_captcha_v2()

    def resolver_2fa(self) -> None:
        self.top.quitar(nm.SELECTORES_2FA[0])
        self.entrar()

    def mostrar_captcha_v2(self) -> None:
        marcar_captcha(self.top, ("anchor",))
        self.top.frames.append(URL_ANCHOR_V2)

    def hay_captcha_visible(self) -> bool:
        return URL_ANCHOR_V2 in self.top.frames

    def resolver_captcha(self, *, entrar: bool = False) -> None:
        quitar_captchas(self.top)
        if entrar:
            self.entrar()

    def entrar(self) -> None:
        self.logueado = True
        self.home()

    def home(self) -> None:
        self.top.reiniciar(url=URL_HOME, titulo="Portal del Consumidor", texto="Hola Ana")
        self.marco.reiniciar()
        self.top.agregar(S("sesion_iniciada"), texto="Cerrar sesión")
        if self.modal_sesion:
            self._modal()

    def _modal(self) -> None:
        def mantener(d: DOMFalso, e: Elemento) -> None:
            self.clicks_mantener += 1
            self.modal_sesion = False
            e.visible = False

        self.top.agregar(S("boton_mantener_sesion"), texto="Permanecer conectado", al_click=mantener)

    def pantalla_no_molestar(self) -> None:
        self.top.reiniciar(url=nm.URL_NO_MOLESTAR, titulo="No Molestar", texto="Registro No Molestar")
        self.top.agregar(S("sesion_iniciada"), texto="Cerrar sesión")
        self.top.agregar(S("acepto_datos"), texto="Acepto que el SERNAC envíe mis datos personales", marcado=self.acepto_marcado)
        self.top.agregar(S("marco_no_molestar"), caja=dict(CAJA_MARCO))
        self.top.frames.append(URL_MARCO_LISTADO)
        if self.modal_sesion:
            self._modal()
        self.listado()

    # ---------------------------------------------------------- iframe: listado
    def listado(self, alerta: str | None = None, bloques: list[str] | None = None, texto_extra: str = "") -> None:
        texto = "Mis solicitudes No Molestar. "
        texto += ("Bloqueos vigentes: " + " | ".join(bloques)) if bloques else "No posee bloqueos."
        if texto_extra:
            texto += " " + texto_extra
        self.marco.reiniciar(url=URL_MARCO_LISTADO, titulo="No Molestar", texto=texto)
        self.marco.agregar(S("boton_nueva_solicitud"), texto="Ingresar nuevo teléfono y/o correo electrónico", al_click=self._nueva)
        if alerta:
            self.marco.agregar(S("confirmacion"), texto=alerta)

    def _nueva(self, dom: DOMFalso, el: Elemento) -> None:
        self.clicks_nueva += 1
        if self.fallos_nueva_solicitud:
            raise self.fallos_nueva_solicitud.pop(0)
        self.formulario()

    # ---------------------------------------------------------- iframe: formulario
    def formulario(self) -> None:
        m = self.marco
        m.reiniciar(url=URL_MARCO_NUEVO, titulo="Nueva solicitud", texto="Ingresar nuevo teléfono y/o correo electrónico")
        m.agregar(S("radio_telefono"), texto="Teléfono", marcado=True, al_marcar=self._radio)
        m.agregar(S("radio_correo"), texto="Correo electrónico", marcado=False, al_marcar=self._radio)
        m.agregar(S("tipo_telefono"), opciones=[("movil", "Móvil"), ("fijo", "Fijo")], valor="movil", etiqueta="Móvil")
        m.agregar(S("campo_telefono"))
        m.agregar(S("campo_correo_usuario"), visible=False)
        m.agregar(S("campo_correo_dominio"), visible=False)
        m.agregar(S("boton_agregar_contacto"), texto="Ingresar", al_click=self._agregar_contacto)
        for t in self.telefonos_precargados:
            m.agregar(S("lista_telefonos"), texto=t)
        for c in self.correos_precargados:
            m.agregar(S("lista_correos"), texto=c)
        m.agregar(S("select_empresa"), visible=False, valor="")
        m.agregar(S("abrir_buscador_empresa"), texto="Ingrese el nombre de la empresa", al_click=self._abrir_select2)
        m.agregar(S("campo_buscar_empresa"), visible=False, al_escribir=self._teclear_busqueda, al_llenar=self._llenar_busqueda)
        m.agregar(S("boton_agregar_empresa"), texto="Ingresar", al_click=self._agregar_empresa)
        m.agregar(S("boton_continuar"), texto="Continuar", al_click=self._continuar)

    def _radio(self, dom: DOMFalso, el: Elemento) -> None:
        correo = el is dom.el(S("radio_correo"))
        dom.el(S("radio_telefono" if correo else "radio_correo")).marcado = False
        for clave in ("tipo_telefono", "campo_telefono"):
            dom.el(S(clave)).visible = not correo
        for clave in ("campo_correo_usuario", "campo_correo_dominio"):
            dom.el(S(clave)).visible = correo

    def _error_contacto(self, dom: DOMFalso, mensaje: str) -> None:
        dom.poner(S("error_contacto"), [Elemento(texto=mensaje)])

    def _agregar_contacto(self, dom: DOMFalso, el: Elemento) -> None:
        dom.quitar(S("error_contacto"))
        if dom.el(S("radio_telefono")).marcado:
            campo = dom.el(S("campo_telefono"))
            numero = "+56" + campo.valor.strip()
            if self.rechazar_telefono or not re.fullmatch(r"\+56[2-9]\d{8}", numero):
                return self._error_contacto(dom, "Ingrese un número de teléfono válido")
            if numero in dom.textos(S("lista_telefonos")):
                return self._error_contacto(dom, "El teléfono ya fue ingresado")
            dom.agregar(S("lista_telefonos"), texto=numero)
            self.contactos_agregados.append(("telefono", numero, dom.el(S("tipo_telefono")).etiqueta))
            campo.valor = ""
        else:
            pre, sub = dom.el(S("campo_correo_usuario")), dom.el(S("campo_correo_dominio"))
            correo = f"{pre.valor.strip()}@{sub.valor.strip()}"
            if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", correo):
                return self._error_contacto(dom, "Ingrese un correo electrónico válido")
            if correo.lower() in [t.lower() for t in dom.textos(S("lista_correos"))]:
                return self._error_contacto(dom, "El correo ya fue ingresado")
            dom.agregar(S("lista_correos"), texto=correo)
            self.contactos_agregados.append(("correo", correo, None))
            pre.valor = sub.valor = ""

    # select2
    def _abrir_select2(self, dom: DOMFalso, el: Elemento) -> None:
        campo = dom.el(S("campo_buscar_empresa"))
        campo.visible = True
        campo.valor = ""
        dom.quitar(S("opciones_empresa"))

    def _teclear_busqueda(self, dom: DOMFalso, el: Elemento, texto: str) -> None:
        self._buscar(dom, el.valor)

    def _llenar_busqueda(self, dom: DOMFalso, el: Elemento, valor: str) -> None:
        if valor:
            self._buscar(dom, valor)
        else:
            dom.quitar(S("opciones_empresa"))

    def _buscar(self, dom: DOMFalso, termino: str) -> None:
        self.busquedas.append(termino)
        resultados = buscar_en_directorio(termino, self.directorio, self.limite_resultados)
        dom.poner(
            S("opciones_empresa"),
            [Elemento(texto=t, datos={"id": i}, al_click=self._elegir) for i, t in resultados],
        )

    def _elegir(self, dom: DOMFalso, el: Elemento) -> None:
        self._seleccionar(dom, el.datos["id"], el.texto)

    def _seleccionar(self, dom: DOMFalso, id_: int, texto: str) -> None:
        sel = dom.el(S("select_empresa"))
        sel.valor, sel.texto = str(id_), texto
        dom.quitar(S("opciones_empresa"))
        dom.el(S("campo_buscar_empresa")).visible = False

    def elegir_manual(self, id_: int) -> None:
        """Simula que la persona elige la opción en el navegador (respuesta «m»)."""
        self._seleccionar(self.marco, id_, dict(self.directorio)[id_])

    def _agregar_empresa(self, dom: DOMFalso, el: Elemento) -> None:
        dom.quitar(S("error_empresa"))
        sel = dom.el(S("select_empresa"))
        if not sel.valor:
            dom.poner(S("error_empresa"), [Elemento(texto="Debe seleccionar una empresa")])
            return
        if any(str(e.datos.get("id")) == sel.valor for e in dom.todos(S("lista_empresas"))):
            dom.poner(S("error_empresa"), [Elemento(texto="La empresa ya fue ingresada")])
            return
        dom.agregar(S("lista_empresas"), texto=sel.texto, datos={"id": int(sel.valor)})
        self.empresas_agregadas.append((int(sel.valor), sel.texto))
        sel.valor, sel.texto = "", ""

    def _continuar(self, dom: DOMFalso, el: Elemento) -> None:
        self.clicks_continuar += 1
        self.eventos.append("continuar")
        tels, mails = dom.textos(S("lista_telefonos")), dom.textos(S("lista_correos"))
        emps = [(e.datos["id"], e.texto) for e in dom.todos(S("lista_empresas"))]
        if not (tels or mails) or not emps:
            return
        self._borrador = {"telefonos": tels, "correos": mails, "empresas": emps}
        self.revisar()

    # ---------------------------------------------------------- iframe: revisar solicitud
    def revisar(self) -> None:
        b, m = self._borrador, self.marco
        m.reiniciar(
            url=URL_MARCO_NUEVO,
            titulo="Nueva solicitud",
            texto="Revisar solicitud. Revise información y seleccione, al menos, un canal a bloquear. "
            + " | ".join(t for _, t in b["empresas"]),
        )
        m.agregar(S("boton_confirmar_telefonos"), texto="Confirmar", al_click=self._confirmar_telefonos)
        k = 0
        for id_, _t in b["empresas"]:
            for tel in b["telefonos"]:
                for canal in ("Llamada", "Sms", "Whatsapp"):
                    m.agregar(S("canales", 0), texto=canal, marcado=k not in self.canales_tel_desmarcados, datos={"empresa": id_, "telefono": tel})
                    k += 1
        k = 0
        for id_, _t in b["empresas"]:
            for correo in b["correos"]:
                m.agregar(S("canales", 1), texto="Bloquear", visible=False, marcado=k not in self.canales_mail_desmarcados, datos={"empresa": id_, "correo": correo})
                k += 1
        m.agregar(S("boton_enviar"), texto="Enviar", visible=False, al_click=self._finalizar)

    def _confirmar_telefonos(self, dom: DOMFalso, el: Elemento) -> None:
        """Solo cambia a la pestaña de correos (no envía nada)."""
        self.eventos.append("confirmar")
        dom.ocultar(S("boton_confirmar_telefonos"), S("canales", 0))
        dom.mostrar(S("canales", 1), S("boton_enviar"))
        if self.captcha_al_enviar:
            self.mostrar_captcha_v2()

    def _finalizar(self, dom: DOMFalso, el: Elemento) -> None:
        self.eventos.append("finalizar")
        self.envios += 1
        b = self._borrador
        self.solicitudes.append(
            {
                "telefonos": list(b["telefonos"]),
                "correos": list(b["correos"]),
                "empresas": list(b["empresas"]),
                "canales_tel": [e.marcado for e in dom.todos(S("canales", 0))],
                "canales_mail": [e.marcado for e in dom.todos(S("canales", 1))],
                "acepto": self.top.el(S("acepto_datos")).marcado if self.top.todos(S("acepto_datos")) else None,
            }
        )
        res = self.resultados_envio[min(self.envios, len(self.resultados_envio)) - 1]
        bloques = [t for _, t in b["empresas"]]
        tipo = res[0]
        if tipo == "ok":
            self.listado(alerta=f"Su solicitud fue ingresada correctamente. N° de solicitud: {res[1]}", bloques=bloques)
        elif tipo == "sin_numero":
            self.listado(alerta="Su solicitud fue ingresada correctamente.", bloques=bloques)
        elif tipo == "listado":
            self.listado(bloques=bloques)
        elif tipo == "folio_en_pagina":
            self.listado(alerta="Solicitud ingresada", bloques=bloques, texto_extra=f"Comprobante — Folio: {res[1]}")
        elif tipo == "nada":
            pass
        elif tipo == "error_click":
            raise PWTimeoutError("Timeout 30000ms exceeded (click en #finalizar)")
        else:  # pragma: no cover - error del test
            raise AssertionError(f"resultado de envío desconocido: {res!r}")


class BuscadorFalso:
    """Doble de BotSernac para resolver_entradas / --mapear: solo métodos públicos de la spec.
    buscar_empresas responde con las búsquedas reales grabadas (por término normalizado)."""

    instancias: list["BuscadorFalso"] = []

    def __init__(self, page: Any = None, persona: Any = None, resultados: dict[str, list[str]] | None = None, **kwargs: Any) -> None:
        base = {k: [t for _, t in v] for k, v in RESULTADOS_WALK.items()}
        self.resultados = {_norm(k): list(v) for k, v in (resultados if resultados is not None else base).items()}
        self.page = page
        self.persona = persona
        self.kwargs = kwargs
        self.busquedas: list[str] = []
        self.logins: list[Any] = []
        self.portales_abiertos = 0
        self.empresas_iniciadas = 0
        BuscadorFalso.instancias.append(self)

    def nueva_empresa(self) -> None:
        self.empresas_iniciadas += 1

    def abrir_portal(self) -> None:
        self.portales_abiertos += 1

    def iniciar_sesion(self, clave: Any) -> None:
        self.logins.append(clave)

    def mantener_sesion(self) -> None:
        return None

    def hay_captcha(self) -> bool:
        return False

    def hay_2fa(self) -> bool:
        return False

    def buscar_empresas(self, termino: str) -> list[str]:
        self.busquedas.append(termino)
        return list(self.resultados.get(_norm(termino), []))

    def enviar_solicitud(self, fila: Any, entradas: Any) -> Any:
        pytest.fail("este flujo nunca debe enviar solicitudes", pytrace=False)


# ------------------------------------------------------------------------- fixtures
def _hash(path: Path) -> str | None:
    return hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else None


@pytest.fixture(scope="session", autouse=True)
def _proteger_archivos_reales():
    """Red de seguridad: los archivos reales del usuario no pueden cambiar durante las pruebas."""
    antes = {p: (_hash(p), p.stat().st_mtime_ns if p.exists() else None) for p in (TRACKER_REAL, TRACKER_TRABAJO, MAPEO_REAL)}
    yield
    for p, (h, mtime) in antes.items():
        ahora = (_hash(p), p.stat().st_mtime_ns if p.exists() else None)
        assert ahora == (h, mtime), f"¡Las pruebas modificaron el archivo real {p}!"


@pytest.fixture(autouse=True)
def _aislar_logging():
    """Restaura los loggers raíz y NOMBRE_LOGGER y cierra handlers de archivo creados por el test
    (en Windows un archivo abierto impide borrar tmp_path)."""
    nombres = [None, nm.NOMBRE_LOGGER]
    estado = {}
    for n in nombres:
        lg = logging.getLogger(n)
        estado[n] = (list(lg.handlers), list(lg.filters), lg.level, lg.propagate, lg.disabled)
    yield
    for n in nombres:
        lg = logging.getLogger(n)
        handlers, filtros, nivel, propagate, disabled = estado[n]
        for h in list(lg.handlers):
            if h in handlers or type(h).__module__.startswith("_pytest"):
                continue
            lg.removeHandler(h)
            try:
                h.close()
            except Exception:
                pass
        lg.filters[:] = filtros
        lg.setLevel(nivel)
        lg.propagate = propagate
        lg.disabled = disabled


@pytest.fixture
def persona() -> nm.Persona:
    return nm.Persona(nombre="Ana Pérez Soto", rut="12345678-5", telefono="912345678", correo=CORREO_PERSONA)


@pytest.fixture
def sleep_falso() -> SleepGrabador:
    return SleepGrabador()


@pytest.fixture
def tracker_real(tmp_path: Path) -> Path:
    """COPIA del tracker real del usuario en tmp_path/real/ (el original nunca se toca)."""
    if not TRACKER_REAL.exists():
        pytest.skip(f"no está {TRACKER_REAL.name} en la carpeta del proyecto")
    destino = tmp_path / "real" / NOMBRE_TRACKER
    destino.parent.mkdir()
    shutil.copy2(TRACKER_REAL, destino)
    return destino


@pytest.fixture
def mapeo_real(tmp_path: Path) -> Path:
    """COPIA del mapeo real del usuario en tmp_path/mapeo/ (el original nunca se toca)."""
    if not MAPEO_REAL.exists():
        pytest.skip("no está mapeo_sernac.csv en la carpeta del proyecto")
    destino = tmp_path / "mapeo" / MAPEO_REAL.name
    destino.parent.mkdir()
    shutil.copy2(MAPEO_REAL, destino)
    return destino


@pytest.fixture
def fabrica_tracker(tmp_path: Path) -> Callable[..., Path]:
    def _crear(filas: list[dict[str, Any]], nombre: str = "tracker.xlsx", **kwargs: Any) -> Path:
        return crear_tracker_real(tmp_path / nombre, filas, **kwargs)

    return _crear


@pytest.fixture
def escribir(tmp_path: Path) -> Callable[..., Path]:
    """escribir(nombre, texto, bom=False) -> Path en tmp_path (UTF-8)."""

    def _escribir(nombre: str, texto: str, bom: bool = False) -> Path:
        p = tmp_path / nombre
        p.write_text(texto, encoding="utf-8-sig" if bom else "utf-8", newline="")
        return p

    return _escribir


@pytest.fixture
def archivos(tmp_path: Path) -> SimpleNamespace:
    """Entradas para main() en tmp_path/entrada (COPIAS de tracker y mapeo reales) y argv base
    (siempre con --salida dentro de tmp_path)."""
    if not (TRACKER_REAL.exists() and MAPEO_REAL.exists()):
        pytest.skip("faltan el tracker real o mapeo_sernac.csv en la carpeta del proyecto")
    entrada = tmp_path / "entrada"
    entrada.mkdir()
    config = entrada / "config.toml"
    config.write_text(CONFIG_VALIDA, encoding="utf-8")
    plantilla = entrada / "plantilla.txt"
    plantilla.write_text(PLANTILLA_VALIDA, encoding="utf-8")
    tracker = entrada / NOMBRE_TRACKER
    shutil.copy2(TRACKER_REAL, tracker)
    mapeo = entrada / MAPEO_REAL.name
    shutil.copy2(MAPEO_REAL, mapeo)
    salida = tmp_path / "salida"

    def argv(*extra: str, tracker_path: Path | None = None, mapeo_path: Path | None = None) -> list[str]:
        return [
            "--tracker", str(tracker_path or tracker),
            "--config", str(config),
            "--plantilla", str(plantilla),
            "--mapeo", str(mapeo_path or mapeo),
            "--salida", str(salida),
            *extra,
        ]

    return SimpleNamespace(config=config, plantilla=plantilla, tracker=tracker, mapeo=mapeo, salida=salida, argv=argv, dir=entrada)


@pytest.fixture
def sin_efectos(monkeypatch: pytest.MonkeyPatch) -> SimpleNamespace:
    """Hace fallar (y registra) getpass, SMTP, navegador e input: para --dry-run y validación."""
    import builtins
    import getpass

    import playwright.sync_api

    registro = SimpleNamespace(llamadas=[])

    def prohibido(nombre: str) -> Callable[..., Any]:
        def _f(*args: Any, **kwargs: Any) -> Any:
            registro.llamadas.append(nombre)
            pytest.fail(f"{nombre} no debía llamarse en esta prueba", pytrace=False)

        return _f

    monkeypatch.setattr(getpass, "getpass", prohibido("getpass.getpass"))
    monkeypatch.setattr(smtplib, "SMTP_SSL", prohibido("smtplib.SMTP_SSL"))
    monkeypatch.setattr(smtplib, "SMTP", prohibido("smtplib.SMTP"))
    monkeypatch.setattr(nm, "lanzar_navegador", prohibido("lanzar_navegador"))
    monkeypatch.setattr(playwright.sync_api, "sync_playwright", prohibido("sync_playwright"))
    monkeypatch.setattr(builtins, "input", prohibido("input"))
    return registro


@pytest.fixture
def tiempos_rapidos(monkeypatch: pytest.MonkeyPatch) -> SleepGrabador:
    """Para main(): pausas con valores fijos reconocibles (empresas 50 s, correos 70 s), resto en 0,
    timeouts cortos y time.sleep falso."""
    import time

    monkeypatch.setitem(nm.TIEMPOS, "pausa_entre_empresas", (50, 50))
    monkeypatch.setitem(nm.TIEMPOS, "pausa_entre_correos", (70, 70))
    monkeypatch.setitem(nm.TIEMPOS, "pausa_entre_acciones", (0, 0))
    monkeypatch.setitem(nm.TIEMPOS, "pausa_entre_busquedas", (0, 0))
    monkeypatch.setitem(nm.TIEMPOS, "backoff_base_s", 0)
    monkeypatch.setitem(nm.TIEMPOS, "retardo_tecleo_ms", 0)
    monkeypatch.setitem(nm.TIEMPOS, "timeout_elemento_ms", 200)
    monkeypatch.setitem(nm.TIEMPOS, "timeout_confirmacion_ms", 200)
    monkeypatch.setitem(nm.TIEMPOS, "timeout_login_s", 1)
    grabador = SleepGrabador()
    monkeypatch.setattr(time, "sleep", grabador)
    return grabador


@pytest.fixture
def nuevo_portal() -> Callable[..., PortalFalso]:
    """nuevo_portal(**opciones de PortalFalso) -> PortalFalso (page en .page)."""
    return lambda **opciones: PortalFalso(**opciones)


@pytest.fixture
def crear_bot(persona: nm.Persona, tmp_path: Path) -> Callable[..., Any]:
    """crear_bot(page, preguntar=..., sleep=..., confirmar=False, ...) -> nm.BotSernac."""

    def _crear(
        page: Any,
        *,
        preguntar: Any = None,
        sleep: Any = None,
        confirmar: bool = False,
        selectores: dict[str, list[str]] | None = None,
        tiempos: dict[str, Any] | None = None,
        capturas_dir: Path | None = None,
        persona_bot: nm.Persona | None = None,
    ) -> Any:
        return nm.BotSernac(
            page,
            persona_bot or persona,
            selectores=selectores or nm.SELECTORES,
            preguntar=preguntar if preguntar is not None else PreguntarGuionado(),
            sleep=sleep if sleep is not None else SleepGrabador(),
            confirmar=confirmar,
            capturas_dir=capturas_dir or (tmp_path / "salida" / "capturas"),
            tiempos=tiempos or tiempos_prueba(),
        )

    return _crear
